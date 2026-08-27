#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""运行高校检索前五步并生成基础阶段输出。"""
import argparse
import json
import re
import sys
import warnings
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


COMMON_ADDRESS_DIR = (
    Path(__file__).resolve().parents[3] / 'components'
)
if str(COMMON_ADDRESS_DIR) not in sys.path:
    sys.path.insert(0, str(COMMON_ADDRESS_DIR))

from normalize_city import normalize_city_name, read_city_prefix_asset
from script_io import write_json_payload, write_workbook_atomically


SCHEMA_VERSION = '1.0'
ASSETS_DIR = Path(__file__).resolve().parents[1] / 'assets'
DEFAULT_SCHOOLS_PATH = ASSETS_DIR / '全国普通高等学校名单.xlsx'
DEFAULT_985_PATH = ASSETS_DIR / '985_universities.xlsx'
DEFAULT_211_PATH = ASSETS_DIR / '211_universities.xlsx'
REQUIRED_COLUMNS = ('序号', '学校名称', '学校标识码', '主管部门', '所在地', '办学层次', '备注')
BASE_HEADER = [
    '序号', '学校名称', '学校标识码', '主管部门',
    '所在地', '办学层次', '院校标签', '办学性质',
]

warnings.filterwarnings(
    'ignore',
    message='Cannot parse header or footer so it will be ignored',
    category=UserWarning,
    module=r'openpyxl\.worksheet\.header_footer',
)

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def _exit_with_error(message):
    """输出统一错误信息并终止脚本。"""
    raise SystemExit(f'错误：{message}')


def _normalize_cell_text(value):
    """把源表单元格转换为去空白字符串。"""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _open_source_workbook(path):
    """以只读模式打开高校源工作簿。"""
    return openpyxl.load_workbook(Path(path), read_only=True, data_only=True)


def _find_header(workbook, alternatives):
    """查找满足任一字段组合的工作表表头。"""
    for sheet in workbook.worksheets:
        for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
            names = [_normalize_cell_text(cell) for cell in row]
            mapping = {name: position for position, name in enumerate(names) if name}
            if any(
                all(column in mapping for column in required)
                for required in alternatives
            ):
                return sheet, row_no, mapping
    return None


def read_moe_school_records(path):
    """读取教育部名单并转换为统一高校字段。"""
    workbook = _open_source_workbook(path)
    try:
        header = _find_header(workbook, (REQUIRED_COLUMNS,))
        if not header:
            raise ValueError(
                '未找到包含全部必填字段的表头：' + '、'.join(REQUIRED_COLUMNS)
            )
        sheet, header_row, columns = header
        rows = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            values = {
                column: _normalize_cell_text(row[columns[column]])
                if len(row) > columns[column] else ''
                for column in REQUIRED_COLUMNS
            }
            school_name = values['学校名称']
            school_identifier = values['学校标识码']
            if not school_name or not school_identifier:
                continue
            rows.append({
                'source_sequence': values['序号'],
                'school_name': school_name,
                'school_identifier': school_identifier,
                'supervising_authority': values['主管部门'],
                'location_city': values['所在地'],
                'education_level': values['办学层次'],
                'source_remark': values['备注'],
            })
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f'教育部名单中未读到有效学校记录：{path}')
    return rows


def read_school_names(path, label):
    """读取院校标签资产中的学校名称。"""
    workbook = _open_source_workbook(path)
    try:
        header = _find_header(
            workbook,
            (('学校名称',), ('school_name',)),
        )
        if not header:
            raise ValueError(f'{label}名单未找到“学校名称”或“school_name”表头：{path}')

        sheet, header_row, columns = header
        column_name = next(
            name for name in ('学校名称', 'school_name') if name in columns
        )
        column = columns[column_name]
        names = set()
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if len(row) <= column:
                continue
            school_name = _normalize_cell_text(row[column])
            if school_name:
                names.add(school_name)
    finally:
        workbook.close()
    if not names:
        raise ValueError(f'{label}名单中未读到学校名称：{path}')
    return names


def filter_schools_by_city(school_records, city):
    """筛选所在地与标准城市名称完全一致的学校。"""
    return [
        school
        for school in school_records
        if school['location_city'] == city
    ]


def classify_school_tag(school, school_names_985, school_names_211):
    """根据办学层次和名单成员关系判定院校标签。"""
    if school['education_level'] != '本科':
        return ''
    if school['school_name'] in school_names_985:
        return '985'
    if school['school_name'] in school_names_211:
        return '211'
    return ''


def classify_school_nature(source_remark):
    """根据教育部源表备注判定办学性质。"""
    if '中外合作办学' in source_remark or '内地与港澳合作办学' in source_remark:
        return '中外合作'
    if '境外高等教育机构' in source_remark:
        return '境外机构'
    if '民办' in source_remark:
        return '民办'
    if not source_remark.strip():
        return '公办'
    return '待核验'


def _validate_school_records(schools, city):
    """校验筛选后的高校记录及派生字段。"""
    if not schools:
        raise ValueError(f'所在地为“{city}”的学校数量为 0')
    school_identifiers = [school['school_identifier'] for school in schools]
    if len(school_identifiers) != len(set(school_identifiers)):
        raise ValueError(f'所在地为“{city}”的学校标识码存在重复')
    if any(school['location_city'] != city for school in schools):
        raise ValueError('筛选结果中存在非目标城市学校')
    if any(
        school['education_level'] != '本科' and school['school_tag']
        for school in schools
    ):
        raise ValueError('非本科院校出现了 985/211 标签')


def _build_base_information_payload(raw_city, source_paths, run_date):
    """构建高校基础阶段数据。"""
    source_schools = read_moe_school_records(source_paths['schools'])
    city = normalize_city_name(raw_city, read_city_prefix_asset())
    city_schools = filter_schools_by_city(source_schools, city)
    school_names_985 = read_school_names(source_paths['985'], '985')
    school_names_211 = read_school_names(source_paths['211'], '211')

    schools = []
    warning_records = []
    for school_record in city_schools:
        school_nature = classify_school_nature(
            school_record['source_remark']
        )
        school = dict(school_record)
        school['school_tag'] = classify_school_tag(
            school_record,
            school_names_985,
            school_names_211,
        )
        school['school_nature'] = school_nature
        if school_nature == '待核验':
            warning_records.append({
                'step': 5,
                'source_sequence': school_record['source_sequence'],
                'school_name': school_record['school_name'],
                'reason': (
                    '备注未匹配办学性质规则：'
                    + school_record['source_remark']
                ),
            })
        schools.append(school)

    _validate_school_records(schools, city)
    return {
        'schema_version': SCHEMA_VERSION,
        'stage': 'base_information',
        'run': {'input_city': raw_city, 'city': city, 'date': run_date},
        'metrics': {
            'school_count': len(schools),
            'warning_count': len(warning_records),
        },
        'schools': schools,
        'warnings': warning_records,
    }


def _build_base_information_workbook(payload):
    """构建保留教育部原序号的基础信息工作簿。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '基础信息'
    sheet.append(BASE_HEADER)
    for school in payload['schools']:
        sheet.append([
            school['source_sequence'],
            school['school_name'],
            school['school_identifier'],
            school['supervising_authority'],
            school['location_city'],
            school['education_level'],
            school['school_tag'],
            school['school_nature'],
        ])

    grid = Side(style='thin', color='B4C7E7')
    cell_border = Border(left=grid, right=grid, top=grid, bottom=grid)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=False, indent=0
        )
        cell.border = cell_border

    table = Table(displayName='BaseInformationTable', ref=f'A1:H{sheet.max_row}')
    table.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium2',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False

    widths = [8, 26, 18, 20, 12, 12, 12, 14]
    for index, width in enumerate(widths, 1):
        column = openpyxl.utils.get_column_letter(index)
        sheet.column_dimensions[column].width = width
    for cell in sheet['C'][1:]:
        cell.number_format = '@'
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=False, indent=0
            )
            cell.border = cell_border
    return workbook


def _validate_base_information_workbook(path, payload):
    """校验基础信息工作簿的结构、类型和对齐格式。"""
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        if workbook.sheetnames != ['基础信息']:
            raise ValueError('基础工作簿只能包含“基础信息”工作表')
        sheet = workbook['基础信息']
        if sheet.max_row - 1 != payload['metrics']['school_count']:
            raise ValueError('基础信息行数与筛选学校数不一致')
        if sheet.max_column != len(BASE_HEADER):
            raise ValueError('基础信息列数不正确')
        if [cell.value for cell in sheet[1]] != BASE_HEADER:
            raise ValueError('基础信息表头不正确')
        school_identifiers = [cell.value for cell in sheet['C'][1:]]
        if any(
            not isinstance(identifier, str) or not identifier.isdigit()
            for identifier in school_identifiers
        ):
            raise ValueError('学校标识码没有按文本保存')
        for row in sheet.iter_rows():
            for cell in row:
                alignment = cell.alignment
                if (
                    alignment.horizontal != 'center'
                    or alignment.vertical != 'center'
                    or bool(alignment.wrap_text)
                    or (alignment.indent or 0) != 0
                ):
                    raise ValueError(f'单元格 {cell.coordinate} 对齐格式不正确')
    finally:
        workbook.close()


def write_base_information_workbook(payload, output_path):
    """原子写出并校验基础信息工作簿。"""
    workbook = _build_base_information_workbook(payload)
    output = write_workbook_atomically(
        workbook,
        output_path,
        lambda path: _validate_base_information_workbook(path, payload),
    )
    return {'output': str(output), 'rows': payload['metrics']['school_count']}


def _validate_source_file(path, label):
    """校验高校私有源文件是否存在。"""
    if not path.is_file():
        _exit_with_error(f'缺少{label}：{path}')


def _build_output_directory_name(city, run_date):
    """生成高校任务输出目录名称。"""
    safe_city = re.sub(r'[<>:"/\\|?*]', '_', city)
    return f'Higher_Education_{safe_city}_{run_date}'


def _build_argument_parser():
    """构建基础阶段命令行参数。"""
    parser = argparse.ArgumentParser(description='运行高校查询前五步基础处理')
    parser.add_argument('--city', required=True, help='用户输入的中国城市名')
    return parser


def main():
    """运行高校基础处理并生成输出。"""
    parser = _build_argument_parser()
    args = parser.parse_args()
    run_date = date.today().isoformat()

    source_paths = {
        'schools': DEFAULT_SCHOOLS_PATH,
        '985': DEFAULT_985_PATH,
        '211': DEFAULT_211_PATH,
    }
    labels = {'schools': '教育部名单', '985': '985名单', '211': '211名单'}
    for key, path in source_paths.items():
        _validate_source_file(path, labels[key])

    try:
        payload = _build_base_information_payload(
            args.city,
            source_paths,
            run_date,
        )
    except (ValueError, OSError) as exc:
        _exit_with_error(str(exc))

    directory_name = _build_output_directory_name(payload['run']['city'], run_date)
    output_dir = Path.cwd().resolve() / directory_name
    output_dir.mkdir(parents=True, exist_ok=True)

    json_path = output_dir / 'base_information.json'
    workbook_path = output_dir / 'base_information.xlsx'
    payload['run']['output_dir'] = str(output_dir)
    payload['outputs'] = {'json': str(json_path), 'workbook': str(workbook_path)}
    try:
        write_base_information_workbook(payload, workbook_path)
        write_json_payload(json_path, payload)
    except (ValueError, OSError) as exc:
        _exit_with_error(str(exc))

    print(json.dumps({
        'stage': payload['stage'],
        'city': payload['run']['city'],
        'date': run_date,
        'school_count': payload['metrics']['school_count'],
        'warning_count': payload['metrics']['warning_count'],
        'output_dir': str(output_dir),
        'workbook': str(workbook_path),
        'json': str(json_path),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
