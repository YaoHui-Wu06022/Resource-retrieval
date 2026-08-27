#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""测试高校最终 Excel 输出。"""

import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl


SCRIPT_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_excel import (  # noqa: E402
    OUTPUT_HEADER,
    SHEET_NAME,
    build_output_rows,
    create_workbook,
    load_processed_records,
    main,
    verify_workbook,
)


def build_address_record(
    source_sequence,
    place_name,
    final_address,
    source_reference,
    campus_name='',
    map_status=None,
):
    """构造测试使用的公共地址记录。"""
    return {
        'place_name': place_name,
        'original_address': final_address,
        'source_nature': 'web_search',
        'source_reference': source_reference,
        'attributes': {
            'source_sequence': str(source_sequence),
            'school_name': place_name.removesuffix(campus_name),
            'school_identifier': '4144010559',
            'supervising_authority': '中央统战部',
            'location_city': '广州市',
            'education_level': '本科',
            'school_tag': '211',
            'school_nature': '公办',
            'campus_name': campus_name,
        },
        'normalized_address': final_address,
        'normalization_status': 'complete' if final_address else 'invalid',
        'normalization_reason': '',
        'map_address': final_address,
        'map_status': map_status or ('consistent' if final_address else 'not_found'),
        'map_reason': '',
        'map_poi_type': '',
        'map_poi_typecode': '',
        'final_address': final_address,
    }


class FinalWorkbookTests(unittest.TestCase):
    """覆盖最终表字段、过滤、排序和显示格式。"""

    def test_output_rows_filter_empty_address_and_renumber(self):
        """空最终地址被剔除并按源序号重新生成展示序号。"""
        address_records = [
            build_address_record(
                2100,
                '后序大学第一校区',
                '广州市天河区示例路2号',
                'https://later.example.edu.cn/',
                '第一校区',
            ),
            build_address_record(
                2000,
                '前序大学',
                '',
                'https://earlier.example.edu.cn/',
            ),
            build_address_record(
                2000,
                '前序大学第二校区',
                '广州市越秀区示例路1号',
                'https://earlier.example.edu.cn/contact',
                '第二校区',
            ),
        ]

        output_rows = build_output_rows(address_records)

        self.assertEqual([row[0] for row in output_rows], [1, 2])
        self.assertEqual(
            [row[1] for row in output_rows],
            ['前序大学第二校区', '后序大学第一校区'],
        )

    def test_same_school_keeps_input_order(self):
        """同一源序号的多条校区记录保持公共输入顺序。"""
        address_records = [
            build_address_record(
                2000,
                '示例大学南校区',
                '广州市海珠区示例路1号',
                'https://example.edu.cn/south',
                '南校区',
            ),
            build_address_record(
                2000,
                '示例大学北校区',
                '广州市白云区示例路2号',
                'https://example.edu.cn/north',
                '北校区',
            ),
        ]

        output_rows = build_output_rows(address_records)

        self.assertEqual(
            [row[1] for row in output_rows],
            ['示例大学南校区', '示例大学北校区'],
        )

    def test_output_rows_show_address_acquisition_method(self):
        """最终行应区分官网提取与地图兜底地址。"""
        output_rows = build_output_rows([
            build_address_record(
                2000,
                '官网大学',
                '广州市天河区示例路1号',
                'https://official.example.edu.cn/',
            ),
            build_address_record(
                2001,
                '兜底大学',
                '广州市白云区示例路2号',
                'https://fallback.example.edu.cn/',
                map_status='fallback',
            ),
        ])

        self.assertEqual([row[7] for row in output_rows], ['官网提取', '地图兜底'])

    def test_workbook_contains_only_confirmed_columns(self):
        """最终工作簿只展示已经确认的九个字段。"""
        workbook = create_workbook([])
        worksheet = workbook[SHEET_NAME]

        self.assertEqual(workbook.sheetnames, [SHEET_NAME])
        self.assertEqual([cell.value for cell in worksheet[1]], OUTPUT_HEADER)
        self.assertNotIn('学校标识码', OUTPUT_HEADER)
        self.assertNotIn('所在地', OUTPUT_HEADER)
        self.assertNotIn('校区名称', OUTPUT_HEADER)

    def test_workbook_has_centered_non_wrapping_cells(self):
        """最终工作簿全部居中且不缩进、不自动换行。"""
        output_rows = build_output_rows([
            build_address_record(
                2000,
                '示例大学中心校区',
                '广州市天河区示例路1号',
                'https://example.edu.cn/contact',
                '中心校区',
            )
        ])
        workbook = create_workbook(output_rows)
        worksheet = workbook[SHEET_NAME]

        for row in worksheet.iter_rows():
            for cell in row:
                self.assertEqual(cell.alignment.horizontal, 'center')
                self.assertEqual(cell.alignment.vertical, 'center')
                self.assertFalse(cell.alignment.wrap_text)
                self.assertIn(cell.alignment.indent, {None, 0, 0.0})
        self.assertIsNotNone(worksheet['I2'].hyperlink)

    def test_main_writes_processed_records(self):
        """命令入口直接把公共处理结果写入最终工作簿。"""
        payload = {
            'schema_version': '1.0',
            'stage': 'processed_address_records',
            'city': '广州市',
            'items': [
                build_address_record(
                    2011,
                    '暨南大学石牌校区',
                    '广州市天河区黄埔大道西601号',
                    'https://www.jnu.edu.cn/contact',
                    '石牌校区',
                )
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_path = root / 'processed_address_records.json'
            output_path = root / '高校建筑查询_广州市_2026-08-26.xlsx'
            input_path.write_text(
                json.dumps(payload, ensure_ascii=False),
                encoding='utf-8',
            )
            arguments = [
                'build_excel.py',
                '--input',
                str(input_path),
                '--output',
                str(output_path),
            ]
            with patch.object(sys, 'argv', arguments), redirect_stdout(io.StringIO()):
                main()
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            worksheet = workbook[SHEET_NAME]
            output_row = [cell.value for cell in worksheet[2]]
            verify_workbook(output_path, 1)

        self.assertEqual(output_row[0], 1)
        self.assertEqual(output_row[1], '暨南大学石牌校区')
        self.assertEqual(output_row[6], '广州市天河区黄埔大道西601号')
        self.assertEqual(output_row[7], '官网提取')
        self.assertEqual(output_row[8], 'https://www.jnu.edu.cn/contact')

    def test_loader_rejects_wrong_stage(self):
        """输入阶段不是公共处理结果时拒绝生成工作簿。"""
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / 'invalid.json'
            input_path.write_text(
                json.dumps({'stage': 'address_records', 'city': '广州市', 'items': []}),
                encoding='utf-8',
            )
            with self.assertRaisesRegex(ValueError, 'processed_address_records'):
                load_processed_records(input_path)


if __name__ == '__main__':
    unittest.main()
