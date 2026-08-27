#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""测试基础教育行政单位表和城市总表生成规则。"""

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


SCRIPT_PATH = Path(__file__).resolve().parents[1] / 'build_excel.py'
SPEC = importlib.util.spec_from_file_location('basic_education_build_excel', SCRIPT_PATH)
BUILD_EXCEL = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = BUILD_EXCEL
SPEC.loader.exec_module(BUILD_EXCEL)


def build_address_record(
    administrative_unit,
    place_name='示例学校',
    final_address='广州市测试区测试路1号',
):
    """构造一条公共地址处理结果。"""
    return {
        'place_name': place_name,
        'source_reference': (
            'https://example.gov.cn/list.xlsx | 学校名录.xlsx '
            '| sheet 学校信息 | row 2 | address 1'
        ),
        'attributes': {
            'administrative_unit': administrative_unit,
            'school_type': '小学',
            'school_nature': '公办',
        },
        'final_address': final_address,
    }


class BuildExcelTests(unittest.TestCase):
    """验证过滤、行政单位约束和工作簿内容。"""

    def test_build_school_output_records_filters_empty_final_address(self):
        """最终地址为空的记录不进入工作簿。"""
        records = BUILD_EXCEL.build_school_output_records(
            '测试区',
            [
                build_address_record('测试区'),
                build_address_record('测试区', final_address=''),
            ],
        )
        self.assertEqual(len(records), 1)

    def test_build_school_output_records_rejects_unit_mismatch(self):
        """记录行政单位必须与所在目录一致。"""
        with self.assertRaisesRegex(ValueError, '与目录不一致'):
            BUILD_EXCEL.build_school_output_records(
                '甲区', [build_address_record('乙区')]
            )

    def test_write_workbook_keeps_fixed_fields_and_hyperlink(self):
        """工作簿保留固定字段、连续序号和来源链接。"""
        records = BUILD_EXCEL.build_school_output_records(
            '测试区', [build_address_record('测试区')]
        )
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_path = Path(temporary_dir) / 'output.xlsx'
            BUILD_EXCEL.write_workbook(
                output_path, [(BUILD_EXCEL.SHEET_NAME, records)]
            )
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            try:
                worksheet = workbook[BUILD_EXCEL.SHEET_NAME]
                self.assertEqual(
                    [cell.value for cell in worksheet[1]],
                    BUILD_EXCEL.OUTPUT_HEADER,
                )
                self.assertEqual(worksheet.cell(2, 1).value, 1)
                self.assertEqual(worksheet.cell(2, 2).value, '测试区')
                self.assertEqual(
                    worksheet.cell(2, 7).value,
                    'https://example.gov.cn/list.xlsx | 学校名录.xlsx | row 2',
                )
                self.assertEqual(
                    worksheet.cell(2, 7).hyperlink.target,
                    'https://example.gov.cn/list.xlsx',
                )
            finally:
                workbook.close()

    def test_city_workbook_contains_summary_and_unit_sheets(self):
        """城市工作簿先放总表，再按行政单位分表。"""
        first_records = BUILD_EXCEL.build_school_output_records(
            '甲区', [build_address_record('甲区', place_name='甲学校')]
        )
        second_records = BUILD_EXCEL.build_school_output_records(
            '乙区', [build_address_record('乙区', place_name='乙学校')]
        )
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_path = Path(temporary_dir) / 'city.xlsx'
            BUILD_EXCEL.write_workbook(output_path, [
                (BUILD_EXCEL.SHEET_NAME, first_records + second_records),
                ('甲区', first_records),
                ('乙区', second_records),
            ])
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [BUILD_EXCEL.SHEET_NAME, '甲区', '乙区'],
                )
                self.assertEqual(workbook['甲区'].cell(2, 1).value, 1)
                self.assertEqual(workbook['乙区'].cell(2, 1).value, 1)
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()
