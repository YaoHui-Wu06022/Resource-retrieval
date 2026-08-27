#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成行政单位级基础教育工作簿并合并为城市总表。"""

import argparse
import json
import re
import sys
from pathlib import Path
from tempfile import NamedTemporaryFile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


SHEET_NAME = '学校信息'
OUTPUT_HEADER = [
    '序号',
    '行政单位',
    '学校名称',
    '学校类型',
    '办学性质',
    '最终地址',
    '信息来源',
]
COLUMN_WIDTHS = [8, 14, 36, 20, 12, 44, 52]
SOURCE_URL_PATTERN = re.compile(r'^https?://[^\s|]+', re.IGNORECASE)
ROW_REFERENCE_PATTERN = re.compile(r'^row\s+\d+$', re.IGNORECASE)


def load_processed_address_records(input_path):
    """读取并校验一个行政单位的公共地址处理结果。"""
    with Path(input_path).resolve().open(encoding='utf-8') as stream:
        address_payload = json.load(stream)
    if not isinstance(address_payload, dict):
        raise ValueError(f'{input_path} 顶层必须是对象')
    if address_payload.get('stage') != 'processed_address_records':
        raise ValueError(f'{input_path} 阶段必须是 processed_address_records')
    if not str(address_payload.get('city') or '').strip():
        raise ValueError(f'{input_path} 缺少 city')
    if not isinstance(address_payload.get('items'), list):
        raise ValueError(f'{input_path} 的 items 必须是数组')
    return address_payload


def collect_administrative_unit_payloads(input_dir):
    """收集各行政单位目录中的公共地址处理结果。"""
    input_root_dir = Path(input_dir).resolve()
    if not input_root_dir.is_dir():
        raise ValueError(f'输入目录不存在：{input_root_dir}')
    administrative_unit_payloads = []
    city_name = ''
    for administrative_unit_dir in sorted(
        input_root_dir.iterdir(), key=lambda unit_path: unit_path.name
    ):
        processed_records_path = (
            administrative_unit_dir / 'processed_address_records.json'
        )
        if not administrative_unit_dir.is_dir() or not processed_records_path.is_file():
            continue
        address_payload = load_processed_address_records(processed_records_path)
        payload_city = str(address_payload['city']).strip()
        if city_name and payload_city != city_name:
            raise ValueError('各行政单位处理结果的 city 不一致')
        city_name = payload_city
        administrative_unit_payloads.append((
            administrative_unit_dir.name,
            administrative_unit_dir,
            address_payload,
        ))
    if not administrative_unit_payloads:
        raise ValueError('没有找到任何 processed_address_records.json')
    return city_name, administrative_unit_payloads


def build_school_output_records(administrative_unit_name, address_records):
    """过滤空最终地址并整理一个行政单位的展示记录。"""
    output_records = []
    for address_record in address_records:
        if not isinstance(address_record, dict):
            raise ValueError(f'{administrative_unit_name}包含非对象地址记录')
        final_address = str(address_record.get('final_address') or '').strip()
        if not final_address:
            continue
        place_name = str(address_record.get('place_name') or '').strip()
        source_reference = str(
            address_record.get('source_reference') or ''
        ).strip()
        record_attributes = address_record.get('attributes') or {}
        if not place_name or not source_reference or not isinstance(record_attributes, dict):
            raise ValueError(
                f'{administrative_unit_name}存在缺少名称、来源或属性的有效记录'
            )
        administrative_unit = str(
            record_attributes.get('administrative_unit') or ''
        ).strip()
        if administrative_unit != administrative_unit_name:
            raise ValueError(
                f'{place_name}的行政单位“{administrative_unit}”与目录不一致'
            )
        output_records.append({
            'administrative_unit': administrative_unit,
            'place_name': place_name,
            'school_type': str(record_attributes.get('school_type') or '').strip(),
            'school_nature': str(record_attributes.get('school_nature') or '').strip(),
            'final_address': final_address,
            'source_reference': source_reference,
        })
    return output_records


def build_worksheet_rows(output_records):
    """为工作簿记录添加连续展示序号。"""
    return [
        [
            sequence,
            record['administrative_unit'],
            record['place_name'],
            record['school_type'],
            record['school_nature'],
            record['final_address'],
            format_source_reference(record['source_reference']),
        ]
        for sequence, record in enumerate(output_records, start=1)
    ]


def extract_source_url(source_reference):
    """从来源定位文本开头提取可点击网页地址。"""
    match = SOURCE_URL_PATTERN.match(str(source_reference or '').strip())
    return match.group(0) if match else ''


def format_source_reference(source_reference):
    """把完整来源定位压缩为网址、文件名和原始行号。"""
    parts = [
        part.strip() for part in str(source_reference or '').split('|')
        if part.strip()
    ]
    source_url = extract_source_url(source_reference)
    file_name = parts[1] if len(parts) > 1 else ''
    row_reference = next(
        (part for part in parts[2:] if ROW_REFERENCE_PATTERN.fullmatch(part)),
        '',
    )
    return ' | '.join(
        part for part in (source_url, file_name, row_reference) if part
    )


def populate_worksheet(worksheet, sheet_name, output_records):
    """向一张工作表写入学校记录并应用固定显示格式。"""
    worksheet.title = sheet_name
    worksheet.append(OUTPUT_HEADER)
    for output_row in build_worksheet_rows(output_records):
        worksheet.append(output_row)

    header_fill = PatternFill('solid', fgColor='4472C4')
    header_font = Font(bold=True, color='FFFFFF')
    thin_side = Side(style='thin', color='D9D9D9')
    cell_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    centered = Alignment(
        horizontal='center',
        vertical='center',
        wrap_text=False,
        indent=0,
    )
    for column_index, column_width in enumerate(COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = (
            column_width
        )
    worksheet.row_dimensions[1].height = 24
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 22
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(OUTPUT_HEADER),
    ):
        for cell in row:
            cell.alignment = centered
            cell.border = cell_border
    for source_cell in worksheet['G'][1:]:
        source_url = extract_source_url(source_cell.value)
        if source_url:
            source_cell.hyperlink = source_url
            source_cell.style = 'Hyperlink'
            source_cell.alignment = centered
            source_cell.border = cell_border
    worksheet.freeze_panes = 'A2'
    worksheet.auto_filter.ref = (
        f'A1:{get_column_letter(len(OUTPUT_HEADER))}{worksheet.max_row}'
    )


def build_workbook(sheet_records):
    """按给定顺序创建总表或行政单位工作表。"""
    if not sheet_records:
        raise ValueError('工作簿至少需要一张工作表')
    workbook = openpyxl.Workbook()
    for index, (sheet_name, output_records) in enumerate(sheet_records):
        worksheet = (
            workbook.active if index == 0 else workbook.create_sheet()
        )
        populate_worksheet(worksheet, sheet_name, output_records)
    return workbook


def verify_worksheet(worksheet, expected_records):
    """验证一张学校工作表的内容、序号和显示格式。"""
    if [cell.value for cell in worksheet[1]] != OUTPUT_HEADER:
        raise ValueError(f'{worksheet.title}工作表字段不正确')
    if worksheet.max_row - 1 != len(expected_records):
        raise ValueError(f'{worksheet.title}工作表记录数不正确')
    for sequence, row_index in enumerate(
        range(2, worksheet.max_row + 1), start=1
    ):
        if worksheet.cell(row_index, 1).value != sequence:
            raise ValueError(f'{worksheet.title}工作表序号不连续')
        for column_index in (2, 3, 6, 7):
            if not str(worksheet.cell(row_index, column_index).value or '').strip():
                raise ValueError(f'{worksheet.title}工作表存在空的必填字段')
        source_cell = worksheet.cell(row_index, 7)
        expected_source = format_source_reference(
            expected_records[row_index - 2]['source_reference']
        )
        if source_cell.value != expected_source:
            raise ValueError(f'{worksheet.title}的信息来源展示格式不正确')
        if not source_cell.hyperlink:
            raise ValueError(f'{worksheet.title}的信息来源没有可点击链接')
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(OUTPUT_HEADER),
    ):
        for cell in row:
            if cell.alignment.horizontal != 'center':
                raise ValueError(f'{worksheet.title}存在未居中的单元格')
            if cell.alignment.vertical != 'center':
                raise ValueError(f'{worksheet.title}存在未垂直居中的单元格')
            if cell.alignment.wrap_text:
                raise ValueError(f'{worksheet.title}不得开启自动换行')
            if cell.alignment.indent not in {None, 0, 0.0}:
                raise ValueError(f'{worksheet.title}不得设置缩进')


def verify_workbook(workbook_path, expected_sheets):
    """重新打开工作簿并验证所有工作表。"""
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        expected_names = [sheet_name for sheet_name, _ in expected_sheets]
        if workbook.sheetnames != expected_names:
            raise ValueError('工作簿的工作表名称或顺序不正确')
        for sheet_name, expected_records in expected_sheets:
            verify_worksheet(workbook[sheet_name], expected_records)
    finally:
        workbook.close()


def write_workbook(output_path, sheet_records):
    """原子保存工作簿并在替换前完成验证。"""
    resolved_output_path = Path(output_path).resolve()
    resolved_output_path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        prefix=f'.{resolved_output_path.stem}.',
        suffix='.xlsx',
        dir=resolved_output_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)
    try:
        workbook = build_workbook(sheet_records)
        try:
            workbook.save(temporary_path)
        finally:
            workbook.close()
        verify_workbook(temporary_path, sheet_records)
        temporary_path.replace(resolved_output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return resolved_output_path


def main():
    """生成所有行政单位工作簿和一个城市总表。"""
    parser = argparse.ArgumentParser(
        description='生成基础教育行政单位表和城市总表'
    )
    parser.add_argument('--input-dir', required=True, help='城市输出目录')
    parser.add_argument('--output', required=True, help='城市总表输出路径')
    arguments = parser.parse_args()
    try:
        city, administrative_unit_payloads = collect_administrative_unit_payloads(
            arguments.input_dir
        )
        merged_records = []
        administrative_unit_outputs = []
        administrative_unit_sheets = []
        for unit_name, unit_dir, address_payload in administrative_unit_payloads:
            unit_records = build_school_output_records(
                unit_name, address_payload['items']
            )
            unit_output_path = write_workbook(
                unit_dir / f'基础教育学校信息_{unit_name}.xlsx',
                [(SHEET_NAME, unit_records)],
            )
            administrative_unit_outputs.append({
                'administrative_unit': unit_name,
                'output': str(unit_output_path),
                'row_count': len(unit_records),
            })
            merged_records.extend(unit_records)
            administrative_unit_sheets.append((unit_name, unit_records))
        city_output = write_workbook(
            arguments.output,
            [(SHEET_NAME, merged_records), *administrative_unit_sheets],
        )
    except (FileNotFoundError, json.JSONDecodeError, ValueError) as exc:
        parser.error(str(exc))
    print(json.dumps({
        'city': city,
        'administrative_unit_count': len(administrative_unit_outputs),
        'row_count': len(merged_records),
        'administrative_unit_outputs': administrative_unit_outputs,
        'output': str(city_output),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
