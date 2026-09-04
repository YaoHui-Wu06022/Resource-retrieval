#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成行政单位级医疗工作簿并合并为城市总表。"""

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from query_city_core.excel_output import (
    ResultWorkbookSpec,
    create_result_workbook,
    main_headers,
    write_result_workbook_atomically,
)
from query_city_core.excel_style import (
    validate_common_worksheet,
    write_workbook_atomically,
)

from medical_scope import (
    abnormal_reason,
    collect_administrative_unit_payloads,
    effective_main_record,
    is_abnormal_record,
    merge_city_records,
    partition_main_records,
)


if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


DOMAIN_HEADERS = ('行政单位', '机构名称', '机构类型', '机构级别')
DOMAIN_WIDTHS = (14, 40, 26, 16)
MAIN_SHEET = '机构信息'
ABNORMAL_SHEET = '异常机构'
EMPTY_LEVEL_VALUES = frozenset({'未定级', '无定级', '无级别'})


def _display_institution_level(value):
    """把无实质等级意义的机构级别显示为空。"""
    level_text = str(value or '').strip()
    return '' if level_text in EMPTY_LEVEL_VALUES else level_text


def _domain_values(record):
    """从记录取机构领域列值。"""
    attributes = record.get('attributes') or {}
    return [
        str(attributes.get('administrative_unit') or '').strip(),
        str(record.get('place_name') or '').strip(),
        str(attributes.get('institution_type') or '').strip(),
        _display_institution_level(attributes.get('institution_level')),
    ]


def build_main_rows(records):
    """过滤无地址记录并返回公共生成器需要的 (领域值, 地址记录)。"""
    rows = []
    for record in records:
        effective = effective_main_record(record)
        if not str(effective.get('final_address') or '').strip():
            continue
        rows.append((_domain_values(record), effective))
    return rows


def build_abnormal_rows(records):
    """构造异常机构行。"""
    return [
        (_domain_values(record), record, abnormal_reason(record))
        for record in records
        if is_abnormal_record(record)
    ]


def _write_workbook(output_path, spec, main_rows, abnormal_rows):
    """用公共生成器原子写出一个医疗工作簿。"""
    return write_result_workbook_atomically(
        spec,
        MAIN_SHEET,
        ABNORMAL_SHEET,
        [(MAIN_SHEET, main_rows)],
        abnormal_rows,
        output_path,
    )


def _validate_main_only_workbook(temporary_path, spec, main_sheets):
    """校验不含异常表的多主表城市总表。"""
    loaded = load_workbook(temporary_path, data_only=False)
    try:
        expected_sheet_names = [sheet_name for sheet_name, _ in main_sheets]
        if loaded.sheetnames != expected_sheet_names:
            raise ValueError('工作簿的工作表名称或顺序不正确')
        headers = main_headers(spec)
        for sheet_name, rows in main_sheets:
            validate_common_worksheet(
                loaded[sheet_name], headers, len(rows)
            )
    finally:
        loaded.close()


def _write_main_only_workbook(output_path, spec, main_sheets):
    """生成不含异常表的城市总表并原子保存。"""
    workbook = create_result_workbook(
        spec,
        MAIN_SHEET,
        ABNORMAL_SHEET,
        main_sheets,
        [],
    )
    workbook.remove(workbook[ABNORMAL_SHEET])

    def validate(temporary_path):
        _validate_main_only_workbook(
            temporary_path, spec, main_sheets
        )

    return write_workbook_atomically(
        workbook, Path(output_path), validate
    )


def main():
    """生成所有行政单位工作簿和一个城市总表。"""
    parser = argparse.ArgumentParser(
        description='生成医疗机构行政单位表和城市总表'
    )
    parser.add_argument(
        '--input-dir',
        required=True,
        help='包含各行政单位目录的 Medical_Institutions 运行目录',
    )
    parser.add_argument(
        '--output',
        required=True,
        help='城市总表输出路径',
    )
    arguments = parser.parse_args()
    try:
        city_name, unit_payloads = collect_administrative_unit_payloads(
            arguments.input_dir
        )
        subdivision_names = [
            str(item.get('name') or '').strip()
            for item in unit_payloads[0][2]['city_context'].get(
                'subdivisions'
            ) or []
        ]
        unit_dir_by_name = {
            unit_name: unit_dir
            for unit_name, unit_dir, _payload in unit_payloads
        }
        unit_records = [
            (record, unit_name)
            for unit_name, _unit_dir, payload in unit_payloads
            for record in payload['items']
        ]
        abnormal_pairs = [
            (record, unit_name)
            for record, unit_name in unit_records
            if is_abnormal_record(record)
        ]
        main_records = [
            record
            for record, _unit_name in unit_records
            if not is_abnormal_record(record)
        ]
        merged_records = merge_city_records(
            main_records, subdivision_names
        )
        main_partitions = partition_main_records(
            merged_records, subdivision_names
        )
        abnormal_partitions = {}
        for record, unit_name in abnormal_pairs:
            abnormal_partitions.setdefault(unit_name, []).append(record)
        spec = ResultWorkbookSpec(DOMAIN_HEADERS, DOMAIN_WIDTHS)
        merged_main_rows = []
        unit_outputs = []
        unit_main_sheets = []
        for unit_name in subdivision_names:
            unit_dir = unit_dir_by_name.get(unit_name)
            if unit_dir is None:
                continue
            main_rows = build_main_rows(
                main_partitions.get(unit_name) or []
            )
            abnormal_rows = build_abnormal_rows(
                abnormal_partitions.get(unit_name) or []
            )
            unit_output_path = _write_workbook(
                unit_dir / f'医疗机构信息_{unit_name}.xlsx',
                spec,
                main_rows,
                abnormal_rows,
            )
            unit_outputs.append({
                'administrative_unit': unit_name,
                'output': str(unit_output_path),
                'row_count': len(main_rows),
                'abnormal_row_count': len(abnormal_rows),
            })
            merged_main_rows.extend(main_rows)
            unit_main_sheets.append((unit_name, main_rows))
        city_main_sheets = [
            (MAIN_SHEET, merged_main_rows),
            *unit_main_sheets,
        ]
        city_output_path = _write_main_only_workbook(
            Path(arguments.output).resolve(),
            spec,
            city_main_sheets,
        )
    except (FileNotFoundError, json.JSONDecodeError, ValueError) as exc:
        parser.error(str(exc))
    print(json.dumps({
        'output': str(city_output_path),
        'city': city_name,
        'administrative_unit_count': len(unit_outputs),
        'row_count': len(merged_main_rows),
        'administrative_unit_outputs': unit_outputs,
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
