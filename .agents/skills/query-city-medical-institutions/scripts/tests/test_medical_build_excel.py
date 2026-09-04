"""医疗机构最终工作簿生成测试。"""

import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

import openpyxl


SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from build_excel import main as build_excel_main  # noqa: E402


def build_city_context(subdivision_names=None):
    """构造城市上下文测试夹具。"""
    subdivision_names = subdivision_names or [
        {'name': '甲区', 'adcode': '1', 'level': 'district'},
    ]
    return {
        'stage': 'city_context',
        'input_city': '示例市',
        'city_name': '示例市',
        'province_name': '示例省',
        'subdivisions': subdivision_names,
    }


def build_record(
    place_name,
    original_address,
    final_address,
    map_match_status='skipped',
    administrative_unit='甲区',
    license_no='',
    institution_level='三级',
):
    """构造一条已处理地址记录。"""
    return {
        'place_name': place_name,
        'original_address': original_address,
        'address_mode': 'government_list',
        'source_nature': 'government_information',
        'source_reference': 'https://example.gov/list | list.html | row 1',
        'attributes': {
            'administrative_unit': administrative_unit,
            'license_administrative_unit': administrative_unit,
            'subdivision_scope': 'subdivision',
            'institution_type': '综合医院',
            'institution_level': institution_level,
            'license_no': license_no,
        },
        'normalized_address': original_address,
        'map_match_status': map_match_status,
        'map_reason': '',
        'final_address': final_address,
        'final_address_source': 'official' if final_address else '',
        'final_address_reason': '政府资料地址' if final_address else '',
    }


def write_unit_payload(unit_dir, city_context, items):
    """写出一个行政单位的 processed_address_records.json。"""
    processed_path = unit_dir / 'processed_address_records.json'
    processed_path.write_text(json.dumps({
        'stage': 'processed_address_records',
        'city_context': city_context,
        'items': items,
        'metrics': {},
    }, ensure_ascii=False), encoding='utf-8')


def run_build_excel(root, output_path):
    """在隔离 argv 下调用最新 build_excel 入口。"""
    old_argv = sys.argv
    sys.argv = [
        'build_excel',
        '--input-dir', str(root),
        '--output', str(output_path),
    ]
    try:
        output = io.StringIO()
        with redirect_stdout(output):
            build_excel_main()
    finally:
        sys.argv = old_argv


class MedicalBuildExcelTests(unittest.TestCase):
    def test_cross_district_rows_route_to_physical_unit_and_dedupe(self):
        """跨目录同址记录去重后进入最终地址所在区工作簿。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first_dir = root / '甲区'
            second_dir = root / '乙区'
            first_dir.mkdir()
            second_dir.mkdir()
            city_context = build_city_context([
                {'name': '甲区', 'adcode': '1', 'level': 'district'},
                {'name': '乙区', 'adcode': '2', 'level': 'district'},
            ])
            first_items = [
                build_record(
                    '某医院乙区门诊部',
                    '示例市乙区水荫路1号',
                    '示例市乙区水荫路1号',
                    administrative_unit='甲区',
                    license_no='A1',
                ),
                build_record(
                    '无地址诊所',
                    '',
                    '',
                ),
            ]
            second_items = [
                build_record(
                    '某医院乙区门诊部',
                    '示例市乙区水荫路1号',
                    '示例市乙区水荫路1号',
                    administrative_unit='乙区',
                    license_no='A1',
                ),
            ]
            write_unit_payload(first_dir, city_context, first_items)
            write_unit_payload(second_dir, city_context, second_items)
            output_path = root / '医疗机构信息_示例市_2026-09-03.xlsx'
            run_build_excel(root, output_path)
            first_workbook = openpyxl.load_workbook(
                first_dir / '医疗机构信息_甲区.xlsx'
            )
            self.assertEqual(first_workbook['机构信息'].max_row, 1)
            self.assertEqual(first_workbook['异常机构'].max_row, 2)
            second_workbook = openpyxl.load_workbook(
                second_dir / '医疗机构信息_乙区.xlsx'
            )
            self.assertEqual(second_workbook['机构信息'].max_row, 2)
            city_workbook = openpyxl.load_workbook(output_path)
            self.assertEqual(
                city_workbook.sheetnames,
                ['机构信息', '甲区', '乙区'],
            )
            self.assertEqual(city_workbook['机构信息'].max_row, 2)

    def test_ungraded_levels_render_blank_through_public_workbook(self):
        """未定级/无定级/无级别经工作簿输出后级别列为空。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            unit_dir = root / '甲区'
            unit_dir.mkdir()
            items = [
                build_record(
                    f'级别示例{index}',
                    f'示例市甲区测试路{index}号',
                    f'示例市甲区测试路{index}号',
                    institution_level=level_value,
                )
                for index, level_value in enumerate(
                    ('未定级', '无定级', '无级别', '三级'),
                    start=1,
                )
            ]
            write_unit_payload(
                unit_dir, build_city_context(), items
            )
            output_path = root / '医疗机构信息_示例市_2026-09-03.xlsx'
            run_build_excel(root, output_path)
            workbook = openpyxl.load_workbook(
                unit_dir / '医疗机构信息_甲区.xlsx'
            )
            level_column = workbook['机构信息']['E']
            self.assertEqual(
                [cell.value or '' for cell in level_column[1:]],
                ['', '', '', '三级'],
            )

    def test_main_builds_unit_and_city_workbooks_without_anomaly_file(self):
        """从各行政单位目录生成区级工作簿与城市总表。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            unit_dir = root / '甲区'
            unit_dir.mkdir()
            items = [
                build_record(
                    '甲区人民医院',
                    '示例市甲区测试路1号',
                    '示例市甲区测试路1号',
                ),
                build_record(
                    '甲区口腔门诊',
                    '示例市甲区测试路2号',
                    '',
                ),
                build_record(
                    '无地址诊所',
                    '',
                    '',
                ),
            ]
            write_unit_payload(
                unit_dir, build_city_context(), items
            )
            output_path = root / '医疗机构信息_示例市_2026-09-03.xlsx'
            run_build_excel(root, output_path)
            self.assertTrue(output_path.is_file())
            city_workbook = openpyxl.load_workbook(output_path)
            self.assertEqual(city_workbook.sheetnames, ['机构信息', '甲区'])
            self.assertNotIn('异常机构', city_workbook.sheetnames)
            main_sheet = city_workbook['机构信息']
            self.assertEqual(main_sheet.max_row, 3)
            unit_sheet = city_workbook['甲区']
            self.assertEqual(unit_sheet.max_row, 3)
            district_path = unit_dir / '医疗机构信息_甲区.xlsx'
            self.assertTrue(district_path.is_file())
            district_workbook = openpyxl.load_workbook(district_path)
            self.assertEqual(
                district_workbook.sheetnames, ['机构信息', '异常机构']
            )
            district_abnormal_sheet = district_workbook['异常机构']
            self.assertEqual(district_abnormal_sheet.max_row, 2)


if __name__ == '__main__':
    unittest.main()
