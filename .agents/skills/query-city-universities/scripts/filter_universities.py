#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从教育部普通高校名单中筛选指定城市的高校。"""

import argparse
import json
import sys
import warnings
from datetime import date
from pathlib import Path

import openpyxl

from query_city_core.address.city import validate_city_context
from query_city_core.excel_style import build_table_workbook
from query_city_core.io_utils import write_json_payload


ASSETS_DIR = Path(__file__).resolve().parents[1] / 'assets'
DEFAULT_SCHOOLS_PATH = ASSETS_DIR / '全国普通高等学校名单.xlsx'
DEFAULT_985_PATH = ASSETS_DIR / '985_universities.xlsx'
DEFAULT_211_PATH = ASSETS_DIR / '211_universities.xlsx'
SOURCE_COLUMNS = (
    '序号', '学校名称', '学校标识码', '主管部门', '所在地', '办学层次', '备注',
)
OUTPUT_COLUMNS = (
    '学校名称', '学校标识码', '主管部门', '所在地', '办学层次', '院校标签', '办学性质',
)
OUTPUT_COLUMN_WIDTHS = (26, 16, 20, 12, 12, 12, 14)


if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

warnings.filterwarnings(
    'ignore',
    message=r'Cannot parse header or footer.*',
    category=UserWarning,
)


def _normalize_cell_text(value):
    """将名单单元格值转换为稳定的文本。"""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _load_workbook_for_reading(path):
    """以只读模式读取来源工作簿，避免模板页眉页脚告警。"""
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def read_city_context(city_context_path):
    """读取城市上下文中的输入城市和标准城市名。"""
    with Path(city_context_path).open(encoding='utf-8') as stream:
        context = json.load(stream)
    return validate_city_context(context)


def _find_header(workbook, required_columns):
    """查找包含全部所需列的工作表及其表头映射。"""
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            header = {
                _normalize_cell_text(value): index
                for index, value in enumerate(row)
                if _normalize_cell_text(value)
            }
            if all(column in header for column in required_columns):
                return sheet, row_number, header
    columns = '、'.join(required_columns)
    raise ValueError(f'未找到包含以下列的工作表：{columns}')


def _read_school_names(path, label):
    """读取 985 或 211 名单中的学校名称。"""
    source_path = Path(path)
    if not source_path.is_file():
        raise FileNotFoundError(f'缺少{label}名单：{source_path}')
    workbook = _load_workbook_for_reading(source_path)
    try:
        try:
            sheet, header_row, header = _find_header(workbook, ('学校名称',))
            name_column = header['学校名称']
        except ValueError:
            sheet, header_row, header = _find_header(workbook, ('school_name',))
            name_column = header['school_name']
        school_names = set()
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_number > header_row:
                school_name = _normalize_cell_text(row[name_column])
                if school_name:
                    school_names.add(school_name)
        return school_names
    finally:
        workbook.close()


def classify_school_tag(education_level, school_name, schools_985, schools_211):
    """返回单列院校标签；985 优先于 211。"""
    if education_level != '本科':
        return ''
    if school_name in schools_985:
        return '985'
    if school_name in schools_211:
        return '211'
    return ''


def classify_school_nature(source_remark):
    """按教育部名单备注归类办学性质。"""
    if '中外合作办学' in source_remark or '内地与港澳合作办学' in source_remark:
        return '中外合作'
    if '境外高等教育机构' in source_remark:
        return '境外机构'
    if '民办' in source_remark:
        return '民办'
    if not source_remark:
        return '公办'
    return '待核验'


def filter_universities(
    city_context_path,
    output_path,
    json_output_path=None,
    schools_path=DEFAULT_SCHOOLS_PATH,
    list_985_path=DEFAULT_985_PATH,
    list_211_path=DEFAULT_211_PATH,
    run_date=None,
):
    """筛选一个城市的高校，并输出固定的七列名单。"""
    city_context = read_city_context(city_context_path)
    city_name = city_context['city_name']
    run_date = run_date or date.today().isoformat()
    source_path = Path(schools_path)
    if not source_path.is_file():
        raise FileNotFoundError(f'缺少教育部名单：{source_path}')

    schools_985 = _read_school_names(list_985_path, '985')
    schools_211 = _read_school_names(list_211_path, '211')
    source_workbook = _load_workbook_for_reading(source_path)
    try:
        source_sheet, header_row, header = _find_header(source_workbook, SOURCE_COLUMNS)
        source_sheet_name = source_sheet.title

        output_rows = []
        schools = []
        warnings = []
        for row_number, row in enumerate(source_sheet.iter_rows(values_only=True), 1):
            if row_number <= header_row:
                continue
            source_record = {
                column: _normalize_cell_text(row[header[column]])
                for column in SOURCE_COLUMNS
            }
            if source_record['所在地'] != city_name:
                continue
            school_tag = classify_school_tag(
                source_record['办学层次'], source_record['学校名称'], schools_985, schools_211,
            )
            school_nature = classify_school_nature(source_record['备注'])
            output_rows.append((
                source_record['学校名称'],
                source_record['学校标识码'],
                source_record['主管部门'],
                source_record['所在地'],
                source_record['办学层次'],
                school_tag,
                school_nature,
            ))
            school = {
                'source_sequence': source_record['序号'],
                'school_name': source_record['学校名称'],
                'school_identifier': source_record['学校标识码'],
                'supervising_authority': source_record['主管部门'],
                'location_city': source_record['所在地'],
                'education_level': source_record['办学层次'],
                'school_tag': school_tag,
                'school_nature': school_nature,
            }
            schools.append(school)
            if school_nature == '待核验':
                warnings.append({
                    'source_sequence': school['source_sequence'],
                    'school_name': school['school_name'],
                    'reason': '备注未匹配办学性质规则：' + source_record['备注'],
                })
    finally:
        source_workbook.close()

    output = Path(output_path)
    output_workbook = build_table_workbook(
        source_sheet_name,
        OUTPUT_COLUMNS,
        output_rows,
        OUTPUT_COLUMN_WIDTHS,
    )
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        output_workbook.save(output)
    finally:
        output_workbook.close()

    json_output = (
        Path(json_output_path)
        if json_output_path
        else output.parent / 'city_universities.json'
    )
    payload = {
        'stage': 'city_universities',
        'city_context': city_context,
        'run': {
            'input_city': city_context['input_city'],
            'city': city_name,
            'date': run_date,
            'output_dir': str(output.parent),
        },
        'metrics': {
            'school_count': len(schools),
            'warning_count': len(warnings),
        },
        'schools': schools,
        'warnings': warnings,
        'outputs': {
            'json': str(json_output),
            'workbook': str(output),
        },
    }
    write_json_payload(json_output, payload)

    return {
        'stage': 'city_universities',
        'city_name': city_name,
        'date': run_date,
        'school_count': len(schools),
        'warning_count': len(warnings),
        'output_dir': str(output.parent),
        'output': str(Path(output_path)),
        'json_output': str(json_output),
    }


def main():
    """运行城市高校名单筛选。"""
    parser = argparse.ArgumentParser(description='筛选城市普通高校名单')
    parser.add_argument('--city-context', required=True)
    args = parser.parse_args()
    try:
        output_dir = Path(args.city_context).resolve().parent
        run_date = date.today().isoformat()
        result = filter_universities(
            args.city_context,
            output_dir / 'city_universities.xlsx',
            output_dir / 'city_universities.json',
            run_date=run_date,
        )
    except (FileNotFoundError, json.JSONDecodeError, OSError, ValueError) as exc:
        parser.error(str(exc))
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
