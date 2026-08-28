"""高校基础阶段入口测试。"""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from prepare_universities import (
    BASE_HEADER,
    DEFAULT_211_PATH,
    DEFAULT_985_PATH,
    DEFAULT_SCHOOLS_PATH,
    _build_argument_parser,
    build_output_directory,
    classify_school_nature,
    classify_school_tag,
    normalize_city_name,
    write_base_information_workbook,
)


class PrepareUniversitiesTest(unittest.TestCase):
    def test_output_directory_uses_project_output_hierarchy(self):
        expected_dir = (
            Path(__file__).resolve().parents[5]
            / 'output'
            / '广州市'
            / '2026-08-27'
            / 'Higher_Education'
        )
        self.assertEqual(
            build_output_directory('广州市', '2026-08-27'),
            expected_dir,
        )

    def test_default_sources_are_bundled_skill_assets(self):
        args = _build_argument_parser().parse_args(['--city', '广州'])
        self.assertFalse(hasattr(args, 'schools'))
        self.assertFalse(hasattr(args, 'list_985'))
        self.assertFalse(hasattr(args, 'list_211'))
        self.assertFalse(hasattr(args, 'date'))
        self.assertFalse(hasattr(args, 'output_root'))
        self.assertFalse(hasattr(args, 'output_dir'))
        self.assertFalse(hasattr(args, 'component_config'))
        self.assertTrue(DEFAULT_SCHOOLS_PATH.is_file())
        self.assertTrue(DEFAULT_985_PATH.is_file())
        self.assertTrue(DEFAULT_211_PATH.is_file())

    def test_city_normalization_uses_exact_location(self):
        self.assertEqual(normalize_city_name.__module__, 'normalize_city')
        locations = {'广州市', '惠州市'}
        self.assertEqual(normalize_city_name('广州', locations), '广州市')
        with self.assertRaises(ValueError):
            normalize_city_name('广', locations)

    def test_nature_classification_priority(self):
        self.assertEqual(
            classify_school_nature('民办中外合作办学'),
            '中外合作',
        )
        self.assertEqual(classify_school_nature('民办'), '民办')
        self.assertEqual(classify_school_nature(''), '公办')

    def test_tag_only_applies_to_undergraduate_schools(self):
        lists = {'测试大学'}
        undergraduate = {
            'school_name': '测试大学',
            'education_level': '本科',
        }
        vocational = {
            'school_name': '测试大学',
            'education_level': '专科',
        }
        self.assertEqual(classify_school_tag(undergraduate, lists, set()), '985')
        self.assertEqual(classify_school_tag(vocational, lists, lists), '')

    def test_workbook_has_eight_centered_columns(self):
        payload = {
            'metrics': {'school_count': 1},
            'schools': [{
                'source_sequence': '2010',
                'school_name': '测试大学',
                'school_identifier': '4144010000',
                'supervising_authority': '测试部门',
                'location_city': '广州市',
                'education_level': '本科',
                'source_remark': '民办',
                'school_tag': '',
                'school_nature': '民办',
            }],
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'base_information.xlsx'
            write_base_information_workbook(payload, output)
            workbook = load_workbook(output)
            try:
                sheet = workbook['基础信息']
                self.assertEqual(sheet.max_column, 8)
                self.assertEqual([cell.value for cell in sheet[1]], BASE_HEADER)
                self.assertEqual(sheet['A2'].value, '2010')
                self.assertNotIn('原始备注', BASE_HEADER)
                for row in sheet.iter_rows(min_row=1, max_row=2):
                    for cell in row:
                        self.assertEqual(cell.alignment.horizontal, 'center')
                        self.assertEqual(cell.alignment.vertical, 'center')
                        self.assertFalse(cell.alignment.wrap_text)
                        self.assertEqual(cell.alignment.indent, 0)
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()
