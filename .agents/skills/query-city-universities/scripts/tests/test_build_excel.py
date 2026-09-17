#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""测试高校最终 Excel 输出。"""

import io
import json
import sys
import tempfile
import unittest
from datetime import date
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl


SCRIPT_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_excel import (  # noqa: E402
    ABNORMAL_HEADERS,
    ABNORMAL_SHEET_NAME,
    OUTPUT_HEADER,
    SHEET_NAME,
    build_abnormal_rows,
    build_output_rows,
    create_workbook,
    load_processed_records,
    main,
    verify_workbook,
)
from university_quality_check import find_output_row_issues  # noqa: E402


def build_address_record(
    source_sequence,
    place_name,
    final_address,
    source_reference,
    campus_name='',
    map_match_status=None,
    final_address_source='',
    map_reason='',
    normalization_reason='',
    school_identifier='4144010559',
):
    """构造测试使用的公共地址记录。"""
    return {
        'place_name': place_name,
        'original_address': final_address,
        'source_nature': 'web_search',
        'source_reference': source_reference,
        'attributes': {
            'source_sequence': str(source_sequence),
            'school_name': place_name.removesuffix(campus_name),
            'school_identifier': school_identifier,
            'supervising_authority': '中央统战部',
            'location_city': '广州市',
            'education_level': '本科',
            'school_tag': '211',
            'school_nature': '公办',
            'campus_name': campus_name,
        },
        'normalized_address': final_address,
        'normalization_status': 'complete' if final_address else 'invalid',
        'normalization_reason': normalization_reason,
        'map_address': final_address,
        'map_match_status': map_match_status or ('consistent' if final_address else 'not_found'),
        'map_reason': map_reason,
        'map_poi_type': '',
        'map_poi_typecode': '',
        'final_address': final_address,
        'final_address_source': final_address_source,
    }


def build_domain_row(record):
    """构造门禁函数可读的展示行。"""
    return ([str(record.get('place_name') or '')], record)


class FinalWorkbookTests(unittest.TestCase):
    """覆盖最终表字段、过滤、排序和显示格式。"""

    def test_abnormal_reason_prefers_attributes_abnormal_reason(self):
        """异常原因优先展示无官网等业务原因。"""
        record = build_address_record(
            2000,
            '无官网示例大学',
            '',
            'https://example.gov.cn/record',
        )
        record['attributes']['abnormal_reason'] = '2026 年新设，未找到独立官网'
        record['map_reason'] = '未找到名称、城市和地址均匹配的 POI'

        abnormal_rows = build_abnormal_rows([record])

        self.assertEqual(len(abnormal_rows), 1)
        self.assertEqual(abnormal_rows[0][2], '2026 年新设，未找到独立官网')

    def test_output_rows_filter_empty_address_and_renumber(self):
        """空最终地址被剔除并按源序号重新生成展示序号。"""
        address_records = [
            build_address_record(
                2100,
                '后序大学第一校区',
                '广州市天河区示例路2号',
                'https://later.example.edu.cn/',
                '第一校区',
            ),
            build_address_record(
                2000,
                '前序大学',
                '',
                'https://earlier.example.edu.cn/',
            ),
            build_address_record(
                2000,
                '前序大学第二校区',
                '广州市越秀区示例路1号',
                'https://earlier.example.edu.cn/contact',
                '第二校区',
            ),
        ]

        output_rows = build_output_rows(address_records, '广州市')

        self.assertEqual(len(output_rows), 2)
        self.assertEqual(
            [row[0][0] for row in output_rows],
            ['前序大学第二校区', '后序大学第一校区'],
        )

    def test_same_school_keeps_input_order(self):
        """同一源序号的多条校区记录保持公共输入顺序。"""
        address_records = [
            build_address_record(
                2000,
                '示例大学南校区',
                '广州市海珠区示例路1号',
                'https://example.edu.cn/south',
                '南校区',
            ),
            build_address_record(
                2000,
                '示例大学北校区',
                '广州市白云区示例路2号',
                'https://example.edu.cn/north',
                '北校区',
            ),
        ]

        output_rows = build_output_rows(address_records, '广州市')

        self.assertEqual(
            [row[0][0] for row in output_rows],
            ['示例大学南校区', '示例大学北校区'],
        )

    def test_output_rows_show_address_acquisition_method(self):
        """最终行应区分官网提取和地图信息。"""
        output_rows = build_output_rows([
            build_address_record(
                2000,
                '官网大学',
                '广州市天河区示例路1号',
                'https://official.example.edu.cn/',
            ),
            build_address_record(
                2001,
                '地图大学',
                '广州市白云区示例路2号',
                'https://map.example.edu.cn/',
                final_address_source='map',
                school_identifier='4144010560',
            ),
        ], '广州市')

        workbook = create_workbook(output_rows)
        worksheet = workbook[SHEET_NAME]
        self.assertEqual(worksheet.cell(2, 9).value, '官网提取')
        self.assertEqual(worksheet.cell(3, 9).value, '地图信息')

    def test_map_same_detail_removes_unlabeled_duplicate(self):
        """地图道路和门牌相同的无校区记录应从最终表中删除。"""
        output_rows = build_output_rows([
            build_address_record(
                2000,
                '示例大学校本部',
                '广州市番禺区小谷围街大学路1号',
                'https://example.edu.cn/campus',
                campus_name='校本部',
            ),
            build_address_record(
                2000,
                '示例大学',
                '广州市番禺区大学路1号',
                'https://example.edu.cn/',
            ),
        ], '广州市')
        self.assertEqual(len(output_rows), 1)
        self.assertEqual(output_rows[0][0][0], '示例大学校本部')

    def test_workbook_contains_only_confirmed_columns(self):
        """最终工作簿只展示已经确认的十一个字段。"""
        workbook = create_workbook([])
        worksheet = workbook[SHEET_NAME]

        self.assertEqual(workbook.sheetnames, [SHEET_NAME, ABNORMAL_SHEET_NAME])
        self.assertEqual([cell.value for cell in worksheet[1]], OUTPUT_HEADER)
        self.assertEqual(OUTPUT_HEADER[7], '地址')
        self.assertNotIn('学校标识码', OUTPUT_HEADER)
        self.assertNotIn('所在地', OUTPUT_HEADER)
        self.assertNotIn('校区名称', OUTPUT_HEADER)

    def test_build_abnormal_rows_lists_empty_address_schools(self):
        """最终地址为空的学校应整理为异常校行并保留原因。"""
        records = [
            build_address_record(
                2000,
                '无地址大学',
                '',
                'https://example.edu.cn/contact',
                map_match_status='not_found',
                map_reason='高德服务正常但未找到结果',
            ),
            build_address_record(
                2001,
                '异地大学',
                '',
                'https://outside.example.edu.cn/',
                map_match_status='skipped',
                normalization_reason='原始地址中的城市与目标城市不一致：深圳市',
                school_identifier='4144010560',
            ),
        ]

        rows = build_abnormal_rows(records)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][0][0], '无地址大学')
        self.assertEqual(rows[0][2], '高德服务正常但未找到结果')
        self.assertEqual(rows[1][0][0], '异地大学')
        self.assertEqual(
            rows[1][2],
            '原始地址中的城市与目标城市不一致：深圳市',
        )

    def test_workbook_contains_abnormal_sheet(self):
        """异常校工作表字段、序号和来源链接正确。"""
        records = [
            build_address_record(
                2000,
                '无地址大学',
                '',
                'https://example.edu.cn/contact',
                map_match_status='not_found',
                map_reason='高德服务正常但未找到结果',
            )
        ]
        workbook = create_workbook([], build_abnormal_rows(records))
        abnormal_sheet = workbook[ABNORMAL_SHEET_NAME]

        self.assertEqual(
            [cell.value for cell in abnormal_sheet[1]],
            ABNORMAL_HEADERS,
        )
        self.assertEqual(abnormal_sheet.cell(2, 1).value, 1)
        self.assertEqual(abnormal_sheet.cell(2, 2).value, '无地址大学')
        self.assertEqual(abnormal_sheet.cell(2, 7).value, '高德服务正常但未找到结果')
        self.assertEqual(
            abnormal_sheet.cell(2, 8).hyperlink.target,
            'https://example.edu.cn/contact',
        )

    def test_workbook_centers_cells_and_wraps_only_address_data(self):
        """最终工作簿全部居中，数据行仅地址列自动换行。"""
        output_rows = build_output_rows([
            build_address_record(
                2000,
                '示例大学中心校区',
                '广州市天河区示例路1号',
                'https://example.edu.cn/contact',
                '中心校区',
            )
        ], '广州市')
        workbook = create_workbook(output_rows)
        worksheet = workbook[SHEET_NAME]

        for row in worksheet.iter_rows():
            for cell in row:
                self.assertEqual(cell.alignment.horizontal, 'center')
                self.assertEqual(cell.alignment.vertical, 'center')
                if cell.row == 1 or cell.column == 8:
                    self.assertTrue(cell.alignment.wrap_text)
                else:
                    self.assertFalse(cell.alignment.wrap_text)
                self.assertIn(cell.alignment.indent, {None, 0, 0.0})
        self.assertIsNone(worksheet.row_dimensions[2].height)
        self.assertIsNotNone(worksheet['K2'].hyperlink)

    def test_main_writes_processed_records(self):
        """命令入口直接把公共处理结果写入最终工作簿。"""
        payload = {
            'stage': 'processed_address_records',
            'city_context': {
                'stage': 'city_context',
                'input_city': '广州市',
                'city_name': '广州市',
                'province_name': '广东省',
                'subdivisions': [],
            },
            'items': [
                build_address_record(
                    2011,
                    '暨南大学石牌校区',
                    '广州市天河区黄埔大道西601号',
                    'https://www.jnu.edu.cn/contact',
                    '石牌校区',
                )
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_path = root / 'processed_address_records.json'
            output_path = root / '高校建筑查询_广州市_2026-08-26.xlsx'
            input_path.write_text(
                json.dumps(payload, ensure_ascii=False),
                encoding='utf-8',
            )
            arguments = [
                'build_excel.py',
                '--input',
                str(input_path),
                '--output',
                str(output_path),
            ]
            with patch.object(sys, 'argv', arguments), redirect_stdout(io.StringIO()):
                main()
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            worksheet = workbook[SHEET_NAME]
            output_row = [cell.value for cell in worksheet[2]]
            verify_workbook(output_path, 1)

        self.assertEqual(output_row[0], 1)
        self.assertEqual(output_row[1], '暨南大学石牌校区')
        self.assertEqual(output_row[6], date.today().isoformat())
        self.assertEqual(output_row[7], '广州市天河区黄埔大道西601号')
        self.assertEqual(output_row[8], '官网提取')
        self.assertEqual(output_row[9], '一致')
        self.assertEqual(output_row[10], 'https://www.jnu.edu.cn/contact')

    def test_loader_rejects_wrong_stage(self):
        """输入阶段不是公共处理结果时拒绝生成工作簿。"""
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / 'invalid.json'
            input_path.write_text(
                json.dumps({
                    'stage': 'address_records',
                    'city_context': {
                        'stage': 'city_context',
                        'input_city': '广州市',
                        'city_name': '广州市',
                        'province_name': '广东省',
                        'subdivisions': [],
                    },
                    'items': [],
                }),
                encoding='utf-8',
            )
            with self.assertRaisesRegex(ValueError, 'processed_address_records'):
                load_processed_records(input_path)


class OutputQualityGateTests(unittest.TestCase):
    """覆盖发布前质量门禁的拦截规则。"""

    def test_contact_label_residue_is_flagged(self):
        """地址残留联系词尾巴时返回问题行。"""
        record = build_address_record(
            3100,
            '示例大学花都校区',
            '广州市花都区工业大道11号 TEL：',
            'https://example.edu.cn/contact',
            '花都校区',
            map_match_status='consistent',
            final_address_source='official',
            school_identifier='4144010861',
        )

        issues = find_output_row_issues([build_domain_row(record)])

        self.assertEqual(len(issues), 1)
        self.assertIn('残留联系词', issues[0])
        self.assertIn('示例大学花都校区', issues[0])

    def test_same_road_no_number_duplicate_is_flagged(self):
        """同校同路已有带门牌行时，无门牌行应被门禁拦截。"""
        numbered = build_address_record(
            3101,
            '示例大学校本部',
            '广州市环市东路465号',
            'https://example.edu.cn/',
            '校本部',
            school_identifier='4144013709',
        )
        incomplete = build_address_record(
            3102,
            '示例大学广州校区',
            '广州市环市东路',
            'https://example.edu.cn/charter',
            '广州校区',
            school_identifier='4144013709',
        )

        issues = find_output_row_issues([
            build_domain_row(numbered),
            build_domain_row(incomplete),
        ])

        self.assertEqual(len(issues), 1)
        self.assertIn('无门牌地址', issues[0])
        self.assertIn('示例大学广州校区', issues[0])

    def test_street_level_official_address_is_not_flagged(self):
        """官网仅公开街道级地址（无门牌）且无同路门牌行时不拦截。"""
        record = build_address_record(
            3103,
            '示例职业技术学院',
            '广州市黄埔区龙湖街道示例职业技术学院',
            'https://example.edu.cn/contact',
            map_match_status='needs_review',
            final_address_source='official',
            school_identifier='4144012575',
        )

        issues = find_output_row_issues([build_domain_row(record)])

        self.assertEqual(issues, [])

    def test_degenerate_road_name_does_not_flag_unrelated_campus(self):
        """中文数字路名解析退化时，不同道路的无门牌校区不得被门禁拦截。"""
        numbered = build_address_record(
            3104,
            '示例职业技术学院北校区',
            '广州市白云区钟落潭镇马沥村广从九路160号',
            'https://example.edu.cn/',
            '北校区',
            school_identifier='4144012743',
        )
        unnumbered = build_address_record(
            3105,
            '示例职业技术学院东校区',
            '广州市天河区龙洞教育园区渔兴路',
            'https://example.edu.cn/campus',
            '东校区',
            school_identifier='4144012743',
        )

        issues = find_output_row_issues([
            build_domain_row(numbered),
            build_domain_row(unnumbered),
        ])

        self.assertEqual(issues, [])


if __name__ == '__main__':
    unittest.main()

