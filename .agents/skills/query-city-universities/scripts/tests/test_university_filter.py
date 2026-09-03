"""城市高校名单筛选测试。"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from university_filter import filter_universities


class FilterUniversitiesTests(unittest.TestCase):
    def write_context(self, path, city_name='广州市'):
        path.write_text(
            json.dumps({
                'stage': 'city_context',
                'input_city': city_name,
                'city_name': city_name,
                'province_name': '广东省',
                'subdivisions': [],
            }, ensure_ascii=False),
            encoding='utf-8',
        )

    def write_source(self, path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '普通高等学校名单'
        sheet.append([
            '序号', '学校名称', '学校标识码', '主管部门', '所在地', '办学层次', '备注',
        ])
        sheet.append([
            1, '广州本科高校', 4144010001, '教育部', '广州市', '本科', '',
        ])
        sheet.append([
            2, '广州民办专科', 4144010002, '广东省教育厅', '广州市', '专科', '民办',
        ])
        sheet.append([
            3, '北京高校', 4144010003, '教育部', '北京市', '本科', '',
        ])
        workbook.save(path)
        workbook.close()

    def write_name_list(self, path, school_names, header='学校名称'):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([header])
        for school_name in school_names:
            sheet.append([school_name])
        workbook.save(path)
        workbook.close()

    def test_filters_and_derives_fixed_output_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            context = root / 'city_context.json'
            source = root / '全国普通高等学校名单.xlsx'
            list_985 = root / '985大学名单.xlsx'
            list_211 = root / '211大学名单.xlsx'
            output = root / 'city_universities.xlsx'
            json_output = root / 'city_universities.json'
            self.write_context(context)
            self.write_source(source)
            self.write_name_list(list_985, ['广州本科高校'], header='school_name')
            self.write_name_list(list_211, ['广州本科高校', '广州民办专科'])

            result = filter_universities(
                context,
                output,
                json_output,
                source,
                list_985,
                list_211,
                '2026-08-28',
            )

            self.assertEqual(result, {
                'stage': 'city_universities',
                'city_name': '广州市',
                'date': '2026-08-28',
                'school_count': 2,
                'warning_count': 0,
                'output_dir': str(root),
                'output': str(output),
                'json_output': str(json_output),
            })
            workbook = load_workbook(output, data_only=True)
            try:
                sheet = workbook['普通高等学校名单']
                self.assertEqual([cell.value for cell in sheet[1]], [
                    '学校名称', '学校标识码', '主管部门', '所在地', '办学层次',
                    '院校标签', '办学性质',
                ])
                self.assertEqual([cell.value for cell in sheet[2]], [
                    '广州本科高校', '4144010001', '教育部', '广州市', '本科',
                    '985', '公办',
                ])
                self.assertEqual([cell.value for cell in sheet[3]], [
                    '广州民办专科', '4144010002', '广东省教育厅', '广州市', '专科',
                    None, '民办',
                ])
            finally:
                workbook.close()
            payload = json.loads(json_output.read_text(encoding='utf-8'))
            self.assertEqual(payload['city_context'], {
                'stage': 'city_context',
                'input_city': '广州市',
                'city_name': '广州市',
                'province_name': '广东省',
                'subdivisions': [],
            })
            self.assertEqual(payload['run'], {
                'input_city': '广州市',
                'city': '广州市',
                'date': '2026-08-28',
                'output_dir': str(root),
            })
            self.assertEqual(payload['outputs'], {
                'json': str(json_output),
                'workbook': str(output),
            })
            self.assertEqual(payload['schools'][0]['source_sequence'], '1')
            self.assertEqual(payload['schools'][0]['school_tag'], '985')

    def test_missing_city_name_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            context = root / 'city_context.json'
            source = root / '全国普通高等学校名单.xlsx'
            list_985 = root / '985大学名单.xlsx'
            list_211 = root / '211大学名单.xlsx'
            context.write_text(json.dumps({
                'stage': 'city_context',
                'input_city': '广州市',
                'province_name': '广东省',
                'subdivisions': [],
            }, ensure_ascii=False), encoding='utf-8')
            self.write_source(source)
            self.write_name_list(list_985, [])
            self.write_name_list(list_211, [])

            with self.assertRaisesRegex(ValueError, 'city_name'):
                filter_universities(
                    context,
                    root / 'city_universities.xlsx',
                    schools_path=source,
                    list_985_path=list_985,
                    list_211_path=list_211,
                )
