#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把高校公共地址处理结果写入最终 Excel。"""

import argparse
from datetime import date
import json
import sys

from query_city_core.city import validate_city_context
from query_city_core.address.common import format_map_match_status
from query_city_core.excel_style import (
    DATE_CELL_PATTERN,
    add_source_hyperlinks,
    build_address_output_values,
    build_table_workbook,
    extend_address_output_columns,
    format_source_reference,
    populate_table_worksheet,
    validate_common_worksheet,
    write_workbook_atomically,
)
from build_university_address import postprocess_university_address_records
from query_city_core.io_utils import read_json_payload


if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


SHEET_NAME = '高校信息'
ABNORMAL_SHEET_NAME = '异常校'
DOMAIN_HEADERS = [
    '序号',
    '学校名称',
    '主管部门',
    '办学层次',
    '院校标签',
    '办学性质',
    '查询日期',
]
COLUMN_WIDTHS = [8, 32, 18, 12, 12, 12, 14]
OUTPUT_HEADER, OUTPUT_WIDTHS = extend_address_output_columns(DOMAIN_HEADERS, COLUMN_WIDTHS)
ABNORMAL_HEADERS = [
    '序号',
    '学校名称',
    '院校标签',
    '办学性质',
    '异常原因',
    '地图匹配状态',
    '信息来源',
    '查询日期',
]
ABNORMAL_WIDTHS = [8, 32, 12, 12, 50, 16, 50, 14]


def load_processed_records(input_path):
    """读取并校验公共层处理完成的高校地址记录。"""
    payload = read_json_payload(input_path)
    if not isinstance(payload, dict):
        raise ValueError('输入文件顶层必须是对象')
    if payload.get('stage') != 'processed_address_records':
        raise ValueError('输入文件阶段必须是 processed_address_records')
    validate_city_context(payload.get('city_context'))
    if not isinstance(payload.get('items'), list):
        raise ValueError('processed_address_records.items 必须是数组')
    return payload


def _read_source_sequence(address_record):
    """读取用于排序的教育部源表序号。"""
    attributes = address_record.get('attributes')
    if not isinstance(attributes, dict):
        raise ValueError('有效地址记录的 attributes 必须是对象')
    source_sequence = str(attributes.get('source_sequence') or '').strip()
    if not source_sequence.isdigit():
        raise ValueError('有效地址记录必须包含纯数字 source_sequence')
    return int(source_sequence)


def build_output_rows(address_records, city_name):
    """过滤无效地址并构造按源表顺序排列的最终表格行。"""
    effective_records = []
    address_records = postprocess_university_address_records(address_records)
    for input_position, address_record in enumerate(address_records):
        if not isinstance(address_record, dict):
            raise ValueError('processed_address_records.items 中的元素必须是对象')
        final_address = str(address_record.get('final_address') or '').strip()
        if not final_address or not final_address.startswith(city_name):
            continue
        place_name = str(address_record.get('place_name') or '').strip()
        source_reference = str(address_record.get('source_reference') or '').strip()
        if not place_name:
            raise ValueError('有效地址记录必须包含非空 place_name')
        if not source_reference:
            raise ValueError(f'{place_name}缺少信息来源')
        source_sequence = _read_source_sequence(address_record)
        effective_records.append((source_sequence, input_position, address_record))

    effective_records.sort(key=lambda entry: (entry[0], entry[1]))
    query_date = date.today().isoformat()
    output_rows = []
    for display_sequence, (_, _, address_record) in enumerate(
        effective_records, start=1
    ):
        attributes = address_record['attributes']
        output_rows.append([
            display_sequence,
            str(address_record['place_name']).strip(),
            str(attributes.get('supervising_authority') or '').strip(),
            str(attributes.get('education_level') or '').strip(),
            str(attributes.get('school_tag') or '').strip(),
            str(attributes.get('school_nature') or '').strip(),
            query_date,
            *build_address_output_values(
                address_record, '官网提取', '地图信息'
            ),
        ])
    return output_rows


def build_abnormal_reason(address_record):
    """从地图、规范化和最终地址原因中取最具体的一项。"""
    attributes = address_record.get('attributes') or {}
    return str(
        (attributes.get('abnormal_reason') or '').strip()
        or address_record.get('map_reason')
        or address_record.get('normalization_reason')
        or address_record.get('final_address_reason')
        or ''
    ).strip()


def build_abnormal_rows(address_records):
    """整理最终地址为空的学校为异常校展示行。"""
    abnormal_rows = []
    for address_record in address_records:
        if not isinstance(address_record, dict):
            raise ValueError('processed_address_records.items 中的元素必须是对象')
        final_address = str(address_record.get('final_address') or '').strip()
        if final_address:
            continue
        place_name = str(address_record.get('place_name') or '').strip()
        source_reference = str(
            address_record.get('source_reference') or ''
        ).strip()
        if not place_name or not source_reference:
            raise ValueError('异常校记录必须包含非空名称和来源')
        attributes = address_record.get('attributes') or {}
        abnormal_rows.append([
            len(abnormal_rows) + 1,
            place_name,
            str(attributes.get('school_tag') or '').strip(),
            str(attributes.get('school_nature') or '').strip(),
            build_abnormal_reason(address_record),
            format_map_match_status(address_record.get('map_match_status')),
            format_source_reference(source_reference),
            date.today().isoformat(),
        ])
    return abnormal_rows


def create_workbook(output_rows, abnormal_rows=()):
    """创建高校信息与异常校两个工作表的最终工作簿。"""
    workbook = build_table_workbook(
        SHEET_NAME,
        OUTPUT_HEADER,
        output_rows,
        OUTPUT_WIDTHS,
    )
    worksheet = workbook[SHEET_NAME]

    add_source_hyperlinks(worksheet, OUTPUT_HEADER.index('信息来源') + 1)

    abnormal_sheet = workbook.create_sheet(ABNORMAL_SHEET_NAME)
    populate_table_worksheet(
        abnormal_sheet,
        ABNORMAL_SHEET_NAME,
        ABNORMAL_HEADERS,
        list(abnormal_rows),
        ABNORMAL_WIDTHS,
    )
    add_source_hyperlinks(
        abnormal_sheet, ABNORMAL_HEADERS.index('信息来源') + 1
    )
    return workbook


def verify_abnormal_worksheet(worksheet, expected_rows):
    """验证异常校工作表的字段、序号和来源链接。"""
    if [cell.value for cell in worksheet[1]] != ABNORMAL_HEADERS:
        raise ValueError('异常校工作表字段不正确')
    if worksheet.max_row - 1 != expected_rows:
        raise ValueError('异常校工作表记录数不正确')
    for sequence, row_index in enumerate(
        range(2, worksheet.max_row + 1), start=1
    ):
        if worksheet.cell(row_index, 1).value != sequence:
            raise ValueError('异常校工作表序号不连续')
        for column_index in (2, 5, 6, 7):
            if not str(worksheet.cell(row_index, column_index).value or '').strip():
                raise ValueError('异常校工作表存在空的必填字段')
        query_date = worksheet.cell(row_index, 8).value
        if not DATE_CELL_PATTERN.fullmatch(str(query_date or '')):
            raise ValueError('异常校工作表存在无效查询日期')
        source_cell = worksheet.cell(row_index, 7)
        if not source_cell.hyperlink:
            raise ValueError('异常校工作表信息来源没有可点击链接')


def verify_workbook(workbook_path, expected_row_count, abnormal_row_count=0):
    """重新读取工作簿并验证字段、数据和基础显示格式。"""
    import openpyxl
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        if workbook.sheetnames != [SHEET_NAME, ABNORMAL_SHEET_NAME]:
            raise ValueError('最终工作簿必须包含高校信息和异常校两个工作表')
        worksheet = workbook[SHEET_NAME]
        validate_common_worksheet(worksheet, OUTPUT_HEADER, expected_row_count)

        for display_sequence, row_index in enumerate(
            range(2, worksheet.max_row + 1), start=1
        ):
            if worksheet.cell(row_index, 1).value != display_sequence:
                raise ValueError('最终工作簿序号不连续')
            if not str(worksheet.cell(row_index, 2).value or '').strip():
                raise ValueError('最终工作簿存在空学校名称')
            acquisition_method = worksheet.cell(row_index, OUTPUT_HEADER.index('地址获取方式') + 1).value
            if acquisition_method not in {'官网提取', '地图信息'}:
                raise ValueError('最终工作簿存在无效地址获取方式')
            query_date = worksheet.cell(row_index, OUTPUT_HEADER.index('查询日期') + 1).value
            if not DATE_CELL_PATTERN.fullmatch(str(query_date or '')):
                raise ValueError('最终工作簿存在无效查询日期')
        verify_abnormal_worksheet(
            workbook[ABNORMAL_SHEET_NAME], abnormal_row_count
        )
    finally:
        workbook.close()


def write_workbook(
    workbook, output_path, expected_row_count, abnormal_row_count=0
):
    """原子保存并复核最终工作簿。"""
    return write_workbook_atomically(
        workbook,
        output_path,
        lambda path: verify_workbook(
            path, expected_row_count, abnormal_row_count
        ),
    )


def main():
    """解析命令行参数并生成最终高校信息工作簿。"""
    parser = argparse.ArgumentParser(description='生成高校最终信息工作簿')
    parser.add_argument('--input', required=True, help='公共层处理完成的JSON')
    parser.add_argument('--output', required=True, help='最终Excel输出路径')
    arguments = parser.parse_args()

    payload = load_processed_records(arguments.input)
    output_rows = build_output_rows(
        payload['items'],
        payload['city_context']['city_name'],
    )
    abnormal_rows = build_abnormal_rows(payload['items'])
    workbook = create_workbook(output_rows, abnormal_rows)
    output = write_workbook(
        workbook,
        arguments.output,
        len(output_rows),
        len(abnormal_rows),
    )
    print(json.dumps({
        'output': str(output),
        'city': payload['city_context']['city_name'],
        'row_count': len(output_rows),
        'abnormal_row_count': len(abnormal_rows),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
