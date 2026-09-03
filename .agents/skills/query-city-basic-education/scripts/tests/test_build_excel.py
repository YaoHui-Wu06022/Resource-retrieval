#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""测试基础教育行政单位表和城市总表生成规则。"""

import importlib.util
import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl


SCRIPT_PATH = Path(__file__).resolve().parents[1] / 'build_excel.py'
SPEC = importlib.util.spec_from_file_location('basic_education_build_excel', SCRIPT_PATH)
BUILD_EXCEL = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = BUILD_EXCEL
SPEC.loader.exec_module(BUILD_EXCEL)


def build_address_record(
    administrative_unit,
    place_name='示例学校',
    final_address='广州市测试区测试路1号',
    map_match_status='skipped',
    final_address_source='',
    publication_date='2026-08-01',
    school_type='小学',
    school_nature='公办',
    source_reference=None,
    map_reason='',
    normalization_reason='',
):
    """构造一条公共地址处理结果。"""
    return {
        'place_name': place_name,
        'source_reference': source_reference or (
            'https://example.gov.cn/list.xlsx | 学校名录.xlsx '
            '| sheet 学校信息 | row 2 | address 1'
        ),
        'attributes': {
            'administrative_unit': administrative_unit,
            'school_type': school_type,
            'school_nature': school_nature,
            'publication_date': publication_date,
        },
        'final_address': final_address,
        'map_match_status': map_match_status,
        'map_reason': map_reason,
        'normalization_reason': normalization_reason,
        'final_address_source': final_address_source,
    }


class BuildExcelTests(unittest.TestCase):
    """验证过滤、行政单位约束和工作簿内容。"""

    def test_build_school_output_records_filters_empty_final_address(self):
        """最终地址为空的记录不进入工作簿。"""
        records = BUILD_EXCEL.build_school_output_records(
            '测试区',
            [
                build_address_record('测试区'),
                build_address_record('测试区', final_address=''),
            ],
        )
        self.assertEqual(len(records), 1)

    def test_build_school_output_records_rejects_unit_mismatch(self):
        """记录行政单位必须与所在目录一致。"""
        with self.assertRaisesRegex(ValueError, '与目录不一致'):
            BUILD_EXCEL.build_school_output_records(
                '甲区', [build_address_record('乙区')]
            )

    def test_build_school_output_records_marks_map_address(self):
        """高德地址记录必须展示对应的地址获取方式。"""
        records = BUILD_EXCEL.build_school_output_records(
            '测试区', [build_address_record('测试区', final_address_source='map')]
        )
        self.assertEqual(records[0]['address_acquisition_method'], '高德地图')

    def test_build_school_output_records_merges_name_variants(self):
        """同址同校的校区别名应合并学校类型。"""
        records = BUILD_EXCEL.build_school_output_records('测试区', [
            build_address_record(
                '测试区',
                place_name='广州市南武中学岭画校区',
                school_type='初中',
                publication_date='2025-05-01',
            ),
            build_address_record(
                '测试区',
                place_name='广州市南武中学（岭南画派纪念校区）',
                school_type='高中',
                publication_date='2026-05-01',
                source_reference=(
                    'https://example.gov.cn/latest.pdf | 最新学校名录.pdf '
                    '| row 8'
                ),
            ),
        ])
        self.assertEqual(len(records), 1)
        self.assertEqual(
            records[0]['place_name'],
            '广州市南武中学（岭南画派纪念校区）',
        )
        self.assertEqual(records[0]['school_type'], '初中、高中')
        self.assertEqual(records[0]['publication_date'], '2026-05-01')
        self.assertIn('latest.pdf', records[0]['source_reference'])

    def test_build_school_output_records_merges_stage_name_variants(self):
        """同址同校的学段名称差异应叠加学校类型。"""
        records = BUILD_EXCEL.build_school_output_records('测试区', [
            build_address_record(
                '测试区',
                place_name='广州龙涛外国语学校（小学）',
                school_type='小学',
            ),
            build_address_record(
                '测试区',
                place_name='广州龙涛外国语学校',
                school_type='完全中学',
            ),
        ])
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]['school_type'], '小学、完全中学')

    def test_build_school_output_records_uses_input_city_name(self):
        """学校身份规范化必须使用输入城市而非固定地名。"""
        records = BUILD_EXCEL.build_school_output_records(
            '西湖区',
            [
                build_address_record(
                    '西湖区',
                    place_name='杭州市西湖区示例中学',
                    final_address='杭州市西湖区示例路1号',
                    school_type='初中',
                ),
                build_address_record(
                    '西湖区',
                    place_name='杭州市示例中学（校本部）',
                    final_address='杭州市西湖区示例路1号',
                    school_type='高中',
                ),
            ],
            city_name='杭州市',
        )
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]['school_type'], '初中、高中')

    def test_build_school_output_records_merges_campus_name_containing_school(self):
        """包含“大学城”等文字的校区别名也应正常融合。"""
        records = BUILD_EXCEL.build_school_output_records('测试区', [
            build_address_record(
                '测试区',
                place_name='示例中学大学城校区',
                school_type='初中',
            ),
            build_address_record(
                '测试区',
                place_name='示例中学大学城新校区',
                school_type='高中',
            ),
        ])
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]['school_type'], '初中、高中')

    def test_build_school_output_records_keeps_different_schools_at_same_address(self):
        """同址但校名主体不同的学校不得合并。"""
        records = BUILD_EXCEL.build_school_output_records('测试区', [
            build_address_record(
                '测试区', place_name='广州市测试区嘉福小学'
            ),
            build_address_record(
                '测试区', place_name='广州市测试区嘉福中学'
            ),
        ])
        self.assertEqual(len(records), 2)

    def test_build_school_output_records_sorts_types_stably(self):
        """行政单位记录按固定类型排序且同类保留来源顺序。"""
        type_names = [
            ('技工学校', '技工院校'),
            ('第一高中', '高中'),
            ('第一幼儿园', '幼儿园'),
            ('第一贯通学校', '小学、初中'),
            ('九年学校', '九年一贯制学校'),
            ('十五年学校', '十五年一贯制学校'),
            ('其他学校', '其他官方类型'),
            ('第二高中', '高中'),
        ]
        records = BUILD_EXCEL.build_school_output_records(
            '测试区',
            [
                build_address_record(
                    '测试区',
                    place_name=place_name,
                    final_address=f'测试市测试区测试路{sequence}号',
                    school_type=school_type,
                )
                for sequence, (place_name, school_type) in enumerate(
                    type_names, start=1
                )
            ],
        )
        self.assertEqual(
            [record['place_name'] for record in records],
            [
                '第一幼儿园',
                '第一贯通学校',
                '九年学校',
                '第一高中',
                '第二高中',
                '十五年学校',
                '技工学校',
                '其他学校',
            ],
        )

    def test_write_workbook_keeps_fixed_fields_and_hyperlink(self):
        """工作簿保留固定字段、连续序号和来源链接。"""
        records = BUILD_EXCEL.build_school_output_records(
            '测试区', [build_address_record('测试区')]
        )
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_path = Path(temporary_dir) / 'output.xlsx'
            BUILD_EXCEL.write_workbook(
                output_path, [(BUILD_EXCEL.SHEET_NAME, records)]
            )
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            try:
                worksheet = workbook[BUILD_EXCEL.SHEET_NAME]
                self.assertEqual(
                    [cell.value for cell in worksheet[1]],
                    BUILD_EXCEL.OUTPUT_HEADER,
                )
                self.assertEqual(BUILD_EXCEL.OUTPUT_HEADER[7], '地址')
                self.assertEqual(
                    worksheet['A1'].fill.fgColor.rgb[-6:], '1F4E78'
                )
                self.assertTrue(worksheet.cell(2, 8).alignment.wrap_text)
                self.assertFalse(worksheet.cell(2, 5).alignment.wrap_text)
                self.assertIsNone(worksheet.row_dimensions[2].height)
                self.assertEqual(worksheet.cell(2, 1).value, 1)
                self.assertEqual(worksheet.cell(2, 2).value, '测试区')
                self.assertEqual(worksheet.cell(2, 9).value, '政府资料')
                self.assertEqual(worksheet.cell(2, 10).value, '未查询')
                self.assertEqual(worksheet.cell(2, 6).value, '2026-08-01')
                self.assertEqual(
                    worksheet.cell(2, 11).value,
                    'https://example.gov.cn/list.xlsx | 学校名录.xlsx | row 2',
                )
                self.assertEqual(
                    worksheet.cell(2, 11).hyperlink.target,
                    'https://example.gov.cn/list.xlsx',
                )
            finally:
                workbook.close()

    def test_workbook_fills_blank_legacy_publication_date(self):
        """旧记录缺少发布日期时应使用当天日期。"""
        records = BUILD_EXCEL.build_school_output_records(
            '测试区', [build_address_record('测试区', publication_date='')]
        )
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_path = Path(temporary_dir) / 'blank-date.xlsx'
            BUILD_EXCEL.write_workbook(
                output_path, [(BUILD_EXCEL.SHEET_NAME, records)]
            )
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            try:
                self.assertEqual(
                    workbook[BUILD_EXCEL.SHEET_NAME].cell(2, 6).value,
                    date.today().isoformat(),
                )
            finally:
                workbook.close()

    def test_build_abnormal_rows_lists_empty_address_schools(self):
        """最终地址为空的学校应整理为异常校行并保留原因。"""
        rows = BUILD_EXCEL.build_abnormal_rows(
            '测试区',
            [
                build_address_record(
                    '测试区',
                    place_name='无地址小学',
                    final_address='',
                    map_match_status='not_found',
                    map_reason='高德服务正常但未找到结果',
                ),
                build_address_record(
                    '测试区',
                    place_name='异地幼儿园',
                    final_address='',
                    map_match_status='skipped',
                    normalization_reason='原始地址中的下级行政区不属于目标城市：天河区',
                ),
            ],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][0][1], '无地址小学')
        self.assertEqual(rows[0][2], '高德服务正常但未找到结果')
        self.assertEqual(rows[1][0][1], '异地幼儿园')
        self.assertIn('不属于目标城市', rows[1][2])

    def test_unit_workbook_contains_abnormal_sheet(self):
        """行政单位工作簿应包含学校信息和异常校两个工作表。"""
        unit_records = BUILD_EXCEL.build_school_output_records(
            '测试区',
            [
                build_address_record('测试区', place_name='正常小学'),
                build_address_record(
                    '测试区',
                    place_name='无地址小学',
                    final_address='',
                    map_match_status='not_found',
                    map_reason='高德服务正常但未找到结果',
                ),
            ],
        )
        abnormal_rows = BUILD_EXCEL.build_abnormal_rows(
            '测试区',
            [
                build_address_record('测试区', place_name='正常小学'),
                build_address_record(
                    '测试区',
                    place_name='无地址小学',
                    final_address='',
                    map_match_status='not_found',
                    map_reason='高德服务正常但未找到结果',
                ),
            ],
        )
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_path = Path(temporary_dir) / 'unit.xlsx'
            BUILD_EXCEL.write_workbook(output_path, [
                (BUILD_EXCEL.SHEET_NAME, unit_records),
                (BUILD_EXCEL.ABNORMAL_SHEET_NAME, abnormal_rows),
            ])
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [BUILD_EXCEL.SHEET_NAME, BUILD_EXCEL.ABNORMAL_SHEET_NAME],
                )
                abnormal_sheet = workbook[BUILD_EXCEL.ABNORMAL_SHEET_NAME]
                self.assertEqual(
                    [cell.value for cell in abnormal_sheet[1]],
                    BUILD_EXCEL.ABNORMAL_HEADERS,
                )
                self.assertEqual(abnormal_sheet.cell(2, 2).value, '测试区')
                self.assertEqual(abnormal_sheet.cell(2, 3).value, '无地址小学')
                self.assertEqual(abnormal_sheet.cell(2, 7).value, '高德服务正常但未找到结果')
            finally:
                workbook.close()

    def test_city_workbook_contains_summary_and_unit_sheets(self):
        """城市工作簿先放总表，再按行政单位分表。"""
        first_records = BUILD_EXCEL.build_school_output_records(
            '甲区', [build_address_record('甲区', place_name='甲学校')]
        )
        second_records = BUILD_EXCEL.build_school_output_records(
            '乙区', [build_address_record('乙区', place_name='乙学校')]
        )
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_path = Path(temporary_dir) / 'city.xlsx'
            BUILD_EXCEL.write_workbook(output_path, [
                (BUILD_EXCEL.SHEET_NAME, first_records + second_records),
                ('甲区', first_records),
                ('乙区', second_records),
            ])
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [BUILD_EXCEL.SHEET_NAME, '甲区', '乙区', BUILD_EXCEL.ABNORMAL_SHEET_NAME],
                )
                self.assertEqual(workbook['甲区'].cell(2, 1).value, 1)
                self.assertEqual(workbook['乙区'].cell(2, 1).value, 1)
                self.assertEqual(
                    workbook[BUILD_EXCEL.SHEET_NAME].cell(2, 3).value,
                    '甲学校',
                )
                self.assertEqual(
                    workbook[BUILD_EXCEL.SHEET_NAME].cell(3, 3).value,
                    '乙学校',
                )
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()

