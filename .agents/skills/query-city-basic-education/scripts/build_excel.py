#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成行政单位级基础教育工作簿并合并为城市总表。"""

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

from query_city_core.city import validate_city_context
from query_city_core.address.common import format_map_match_status
from query_city_core.io_utils import read_json_payload
from query_city_core.excel_style import (
    DATE_CELL_PATTERN,
    add_source_hyperlinks,
    build_address_output_values,
    extend_address_output_columns,
    format_date_cell,
    format_source_reference,
    populate_table_worksheet,
    validate_common_worksheet,
    write_workbook_atomically,
)
from normalize_school_records import merge_school_types


if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


SHEET_NAME = '学校信息'
ABNORMAL_SHEET_NAME = '异常校'
DOMAIN_HEADERS = [
    '序号',
    '行政单位',
    '学校名称',
    '学校类型',
    '办学性质',
    '发布日期',
]
DOMAIN_WIDTHS = [8, 14, 36, 20, 12, 18]
OUTPUT_HEADER, COLUMN_WIDTHS = extend_address_output_columns(DOMAIN_HEADERS, DOMAIN_WIDTHS)
ABNORMAL_HEADERS = [
    '序号',
    '行政单位',
    '学校名称',
    '学校类型',
    '办学性质',
    '异常原因',
    '地图匹配状态',
    '信息来源',
    '发布日期',
]
ABNORMAL_WIDTHS = [8, 14, 36, 20, 12, 50, 16, 50, 18]
SCHOOL_NAME_SEPARATOR_PATTERN = re.compile(r'[\s（）()【】\[\]·•]+')
SCHOOL_CAMPUS_SUFFIX_PATTERN = re.compile(
    r'(?P<base>.+?(?:幼儿园|小学|中学|学校))'
    r'(?:校本部|本部|(?:(?!(?:幼儿园|小学|中学|学校)).){1,20}'
    r'(?:校区|园区))$'
)
SCHOOL_STAGE_SUFFIX_PATTERN = re.compile(
    r'(?:(?<=学校)(?:小学|初中|高中)|(?:小学|初中|高中)部)$'
)


SCHOOL_TYPE_SORT_PRIORITIES = {
    '幼儿园': 0,
    '小学': 1,
    '九年一贯制学校': 2,
    '初中': 3,
    '完全中学': 4,
    '高中': 5,
    '十二年一贯制学校': 6,
    '特殊教育学校': 8,
    '中等职业学校': 9,
    '职业高级中学': 10,
    '技工院校': 11,
    '专门学校': 12,
}
OTHER_CONSISTENT_SCHOOL_PRIORITY = 7
OTHER_SCHOOL_TYPE_PRIORITY = 13


def load_processed_address_records(input_path):
    """读取并校验一个行政单位的公共地址处理结果。"""
    address_payload = read_json_payload(input_path)
    if not isinstance(address_payload, dict):
        raise ValueError(f'{input_path} 顶层必须是对象')
    if address_payload.get('stage') != 'processed_address_records':
        raise ValueError(f'{input_path} 阶段必须是 processed_address_records')
    try:
        validate_city_context(address_payload.get('city_context'))
    except ValueError as exc:
        raise ValueError(f'{input_path} 的城市上下文无效：{exc}') from exc
    if not isinstance(address_payload.get('items'), list):
        raise ValueError(f'{input_path} 的 items 必须是数组')
    return address_payload


def collect_administrative_unit_payloads(input_dir):
    """收集各行政单位目录中的公共地址处理结果。"""
    input_root_dir = Path(input_dir).resolve()
    if not input_root_dir.is_dir():
        raise ValueError(f'输入目录不存在：{input_root_dir}')
    administrative_unit_payloads = []
    city_name = ''
    for administrative_unit_dir in sorted(
        input_root_dir.iterdir(), key=lambda unit_path: unit_path.name
    ):
        processed_records_path = (
            administrative_unit_dir / 'processed_address_records.json'
        )
        if not administrative_unit_dir.is_dir() or not processed_records_path.is_file():
            continue
        address_payload = load_processed_address_records(processed_records_path)
        payload_city = address_payload['city_context']['city_name']
        if city_name and payload_city != city_name:
            raise ValueError('各行政单位处理结果的 city 不一致')
        city_name = payload_city
        administrative_unit_payloads.append((
            administrative_unit_dir.name,
            administrative_unit_dir,
            address_payload,
        ))
    if not administrative_unit_payloads:
        raise ValueError('没有找到任何 processed_address_records.json')
    return city_name, administrative_unit_payloads


def normalize_school_identity(place_name, administrative_unit, city_name):
    """规范学校名称以识别同址校区别名和学段写法。"""
    school_identity = SCHOOL_NAME_SEPARATOR_PATTERN.sub('', place_name)
    if city_name and school_identity.startswith(city_name):
        school_identity = school_identity[len(city_name):]
    if administrative_unit and school_identity.startswith(administrative_unit):
        school_identity = school_identity[len(administrative_unit):]
    school_identity = SCHOOL_STAGE_SUFFIX_PATTERN.sub('', school_identity)
    campus_match = SCHOOL_CAMPUS_SUFFIX_PATTERN.fullmatch(school_identity)
    return campus_match.group('base') if campus_match else school_identity


def deduplicate_school_output_records(output_records, city_name):
    """按学校身份和最终地址合并同址别名记录。"""
    unique_records = []
    record_indexes = {}
    for output_record in output_records:
        record_key = (
            output_record['administrative_unit'],
            output_record['final_address'],
            normalize_school_identity(
                output_record['place_name'],
                output_record['administrative_unit'],
                city_name,
            ),
        )
        if record_key not in record_indexes:
            record_indexes[record_key] = len(unique_records)
            unique_records.append(output_record)
            continue
        record_index = record_indexes[record_key]
        current_record = unique_records[record_index]
        school_type = merge_school_types(
            current_record['school_type'], output_record['school_type']
        )
        school_nature = (
            current_record['school_nature'] or output_record['school_nature']
        )
        if output_record['publication_date'] > current_record['publication_date']:
            current_record = output_record.copy()
            unique_records[record_index] = current_record
        current_record['school_type'] = school_type
        current_record['school_nature'] = school_nature
    return unique_records


def resolve_school_type_sort_priority(school_type):
    """确定学校类型在行政单位表中的排序优先级。"""
    priorities = []
    for type_name in str(school_type or '').split('、'):
        if type_name in SCHOOL_TYPE_SORT_PRIORITIES:
            priorities.append(SCHOOL_TYPE_SORT_PRIORITIES[type_name])
        elif re.fullmatch(r'.+年一贯制学校', type_name):
            priorities.append(OTHER_CONSISTENT_SCHOOL_PRIORITY)
        else:
            priorities.append(OTHER_SCHOOL_TYPE_PRIORITY)
    return min(priorities, default=OTHER_SCHOOL_TYPE_PRIORITY)


def sort_school_output_records(output_records):
    """按学校类型稳定排序一个行政单位的展示记录。"""
    return sorted(
        output_records,
        key=lambda record: resolve_school_type_sort_priority(
            record['school_type']
        ),
    )


def build_school_output_records(
    administrative_unit_name, address_records, city_name=''
):
    """过滤空最终地址并整理一个行政单位的展示记录。"""
    output_records = []
    for address_record in address_records:
        if not isinstance(address_record, dict):
            raise ValueError(f'{administrative_unit_name}包含非对象地址记录')
        final_address = str(address_record.get('final_address') or '').strip()
        if not final_address:
            continue
        place_name = str(address_record.get('place_name') or '').strip()
        source_reference = str(
            address_record.get('source_reference') or ''
        ).strip()
        record_attributes = address_record.get('attributes') or {}
        if not place_name or not source_reference or not isinstance(record_attributes, dict):
            raise ValueError(
                f'{administrative_unit_name}存在缺少名称、来源或属性的有效记录'
            )
        administrative_unit = str(
            record_attributes.get('administrative_unit') or ''
        ).strip()
        if administrative_unit != administrative_unit_name:
            raise ValueError(
                f'{place_name}的行政单位“{administrative_unit}”与目录不一致'
            )
        output_records.append({
            'administrative_unit': administrative_unit,
            'place_name': place_name,
            'school_type': str(record_attributes.get('school_type') or '').strip(),
            'school_nature': str(record_attributes.get('school_nature') or '').strip(),
            'publication_date': format_date_cell(
                record_attributes.get('publication_date'),
                date.today().isoformat(),
            ),
            'final_address': final_address,
            'address_acquisition_method': (
                '高德地图'
                if address_record.get('final_address_source') == 'map'
                else '政府资料'
            ),
            'map_match_status': address_record.get('map_match_status') or 'skipped',
            'source_reference': source_reference,
        })
    return sort_school_output_records(
        deduplicate_school_output_records(output_records, city_name)
    )


def build_worksheet_rows(output_records):
    """为工作簿记录添加连续展示序号。"""
    return [
        [
            sequence,
            record['administrative_unit'],
            record['place_name'],
            record['school_type'],
            record['school_nature'],
            record['publication_date'],
            *build_address_output_values(
                {
                    'final_address': record['final_address'],
                    'final_address_source': 'map' if record['address_acquisition_method'] == '高德地图' else 'official',
                    'map_match_status': record['map_match_status'],
                    'source_reference': record['source_reference'],
                },
                '政府资料', '高德地图',
            ),
        ]
        for sequence, record in enumerate(output_records, start=1)
    ]


def build_abnormal_reason(address_record):
    """从地图、规范化和最终地址原因中取最具体的一项。"""
    return str(
        address_record.get('map_reason')
        or address_record.get('normalization_reason')
        or address_record.get('final_address_reason')
        or ''
    ).strip()


def build_abnormal_rows(administrative_unit_name, address_records):
    """整理行政单位内最终地址为空的学校为异常校展示行。"""
    abnormal_rows = []
    for address_record in address_records:
        if not isinstance(address_record, dict):
            raise ValueError(f'{administrative_unit_name}包含非对象地址记录')
        final_address = str(address_record.get('final_address') or '').strip()
        if final_address:
            continue
        place_name = str(address_record.get('place_name') or '').strip()
        source_reference = str(
            address_record.get('source_reference') or ''
        ).strip()
        if not place_name or not source_reference:
            raise ValueError(
                f'{administrative_unit_name}存在缺少名称或来源的异常记录'
            )
        record_attributes = address_record.get('attributes') or {}
        abnormal_rows.append([
            len(abnormal_rows) + 1,
            administrative_unit_name,
            place_name,
            str(record_attributes.get('school_type') or '').strip(),
            str(record_attributes.get('school_nature') or '').strip(),
            build_abnormal_reason(address_record),
            format_map_match_status(address_record.get('map_match_status')),
            format_source_reference(source_reference),
            format_date_cell(
                record_attributes.get('publication_date'),
                date.today().isoformat(),
            ),
        ])
    return abnormal_rows


def populate_worksheet(worksheet, sheet_name, output_records):
    """使用公共样式写入一张学校工作表并添加来源链接。"""
    populate_table_worksheet(
        worksheet,
        sheet_name,
        OUTPUT_HEADER,
        build_worksheet_rows(output_records),
        COLUMN_WIDTHS,
    )
    add_source_hyperlinks(worksheet, OUTPUT_HEADER.index('信息来源') + 1)


def populate_abnormal_worksheet(worksheet, sheet_name, abnormal_rows):
    """使用公共样式写入异常校工作表并添加来源链接。"""
    populate_table_worksheet(
        worksheet,
        sheet_name,
        ABNORMAL_HEADERS,
        abnormal_rows,
        ABNORMAL_WIDTHS,
    )
    add_source_hyperlinks(worksheet, ABNORMAL_HEADERS.index('信息来源') + 1)


def build_workbook(sheet_records):
    """按给定顺序创建总表或行政单位工作表。"""
    if not sheet_records:
        raise ValueError('工作簿至少需要一张工作表')
    workbook = openpyxl.Workbook()
    for index, (sheet_name, output_records) in enumerate(sheet_records):
        worksheet = (
            workbook.active if index == 0 else workbook.create_sheet()
        )
        if sheet_name == ABNORMAL_SHEET_NAME:
            populate_abnormal_worksheet(worksheet, sheet_name, output_records)
        else:
            populate_worksheet(worksheet, sheet_name, output_records)
    return workbook


def verify_worksheet(worksheet, expected_records):
    """验证一张学校工作表的内容、序号和显示格式。"""
    validate_common_worksheet(worksheet, OUTPUT_HEADER, len(expected_records))
    for sequence, row_index in enumerate(
        range(2, worksheet.max_row + 1), start=1
    ):
        if worksheet.cell(row_index, 1).value != sequence:
            raise ValueError(f'{worksheet.title}工作表序号不连续')
        for column_index in (2, 3, 7, 8, 9, 10):
            if not str(worksheet.cell(row_index, column_index).value or '').strip():
                raise ValueError(f'{worksheet.title}工作表存在空的必填字段')
        acquisition_method = worksheet.cell(row_index, 8).value
        if acquisition_method not in {'政府资料', '高德地图'}:
            raise ValueError(f'{worksheet.title}的地址获取方式不正确')
        source_cell = worksheet.cell(row_index, 10)
        expected_source = format_source_reference(
            expected_records[row_index - 2]['source_reference']
        )
        if source_cell.value != expected_source:
            raise ValueError(f'{worksheet.title}的信息来源展示格式不正确')
        if not source_cell.hyperlink:
            raise ValueError(f'{worksheet.title}的信息来源没有可点击链接')
    for row_index, row in enumerate(worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(OUTPUT_HEADER),
    ), start=1):
        for column_index, cell in enumerate(row, start=1):
            if cell.alignment.horizontal != 'center':
                raise ValueError(f'{worksheet.title}存在未居中的单元格')
            if cell.alignment.vertical != 'center':
                raise ValueError(f'{worksheet.title}存在未垂直居中的单元格')
            expected_wrap = row_index == 1 or column_index == 7
            if bool(cell.alignment.wrap_text) != expected_wrap:
                raise ValueError(f'{worksheet.title}的自动换行格式不正确')
            if cell.alignment.indent not in {None, 0, 0.0}:
                raise ValueError(f'{worksheet.title}不得设置缩进')


def verify_workbook(workbook_path, expected_sheets):
    """重新打开工作簿并验证所有工作表。"""
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        expected_names = [sheet_name for sheet_name, _ in expected_sheets]
        if workbook.sheetnames != expected_names:
            raise ValueError('工作簿的工作表名称或顺序不正确')
        for sheet_name, expected_records in expected_sheets:
            if sheet_name == ABNORMAL_SHEET_NAME:
                verify_abnormal_worksheet(
                    workbook[sheet_name], len(expected_records)
                )
            else:
                verify_worksheet(workbook[sheet_name], expected_records)
    finally:
        workbook.close()


def verify_abnormal_worksheet(worksheet, expected_rows):
    """验证异常校工作表的字段、序号和来源链接。"""
    if [cell.value for cell in worksheet[1]] != ABNORMAL_HEADERS:
        raise ValueError('异常校工作表字段不正确')
    if worksheet.max_row - 1 != expected_rows:
        raise ValueError('异常校工作表记录数不正确')
    for sequence, row_index in enumerate(
        range(2, worksheet.max_row + 1), start=1
    ):
        if worksheet.cell(row_index, 1).value != sequence:
            raise ValueError('异常校工作表序号不连续')
        for column_index in (2, 3, 6, 7, 8):
            if not str(worksheet.cell(row_index, column_index).value or '').strip():
                raise ValueError('异常校工作表存在空的必填字段')
        publication_date = worksheet.cell(row_index, 9).value
        if not DATE_CELL_PATTERN.fullmatch(str(publication_date or '')):
            raise ValueError('异常校工作表存在无效发布日期')
        source_cell = worksheet.cell(row_index, 8)
        if not source_cell.hyperlink:
            raise ValueError('异常校工作表信息来源没有可点击链接')


def write_workbook(output_path, sheet_records):
    """原子保存工作簿并在替换前完成验证。"""
    workbook = build_workbook(sheet_records)
    return write_workbook_atomically(
        workbook,
        output_path,
        lambda temporary_path: verify_workbook(temporary_path, sheet_records),
    )


def main():
    """生成所有行政单位工作簿和一个城市总表。"""
    parser = argparse.ArgumentParser(
        description='生成基础教育行政单位表和城市总表'
    )
    parser.add_argument('--input-dir', required=True, help='城市输出目录')
    parser.add_argument('--output', required=True, help='城市总表输出路径')
    arguments = parser.parse_args()
    try:
        city, administrative_unit_payloads = collect_administrative_unit_payloads(
            arguments.input_dir
        )
        merged_records = []
        administrative_unit_outputs = []
        administrative_unit_sheets = []
        for unit_name, unit_dir, address_payload in administrative_unit_payloads:
            unit_records = build_school_output_records(
                unit_name, address_payload['items'], city
            )
            unit_abnormal_rows = build_abnormal_rows(
                unit_name, address_payload['items']
            )
            unit_output_path = write_workbook(
                unit_dir / f'基础教育学校信息_{unit_name}.xlsx',
                [
                    (SHEET_NAME, unit_records),
                    (ABNORMAL_SHEET_NAME, unit_abnormal_rows),
                ],
            )
            administrative_unit_outputs.append({
                'administrative_unit': unit_name,
                'output': str(unit_output_path),
                'row_count': len(unit_records),
                'abnormal_row_count': len(unit_abnormal_rows),
            })
            merged_records.extend(unit_records)
            administrative_unit_sheets.append((unit_name, unit_records))
        city_output = write_workbook(
            arguments.output,
            [(SHEET_NAME, merged_records), *administrative_unit_sheets],
        )
    except (FileNotFoundError, json.JSONDecodeError, ValueError) as exc:
        parser.error(str(exc))
    print(json.dumps({
        'city': city,
        'administrative_unit_count': len(administrative_unit_outputs),
        'row_count': len(merged_records),
        'administrative_unit_outputs': administrative_unit_outputs,
        'output': str(city_output),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
