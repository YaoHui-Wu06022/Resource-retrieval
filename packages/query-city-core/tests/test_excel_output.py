"""统一地址结果工作簿生成器测试。"""

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from query_city_core.excel_output import (
    ResultWorkbookSpec,
    abnormal_column_widths,
    abnormal_headers,
    create_result_workbook,
    main_column_widths,
    main_headers,
    verify_result_workbook,
    write_result_workbook_atomically,
)


DOMAIN_HEADERS = ('区域', '对象名称', '对象类型')
DOMAIN_WIDTHS = (14, 36, 20)
QUERY_DATE = '2026-09-03'


def build_spec(query_date=QUERY_DATE):
    return ResultWorkbookSpec(
        DOMAIN_HEADERS,
        DOMAIN_WIDTHS,
        query_date=query_date,
    )


def build_address_record(
    *,
    final_address='示例大道1号',
    source='official',
    map_status='consistent',
    source_reference='https://www.example.gov.cn/list.html | 名录.xlsx | row 3',
):
    return {
        'final_address': final_address,
        'final_address_source': source,
        'map_match_status': map_status,
        'source_reference': source_reference,
    }


class HeaderTests(unittest.TestCase):
    def test_fixed_headers_put_domain_columns_between_sequence_and_common_columns(self):
        spec = build_spec()
        self.assertEqual(
            main_headers(spec),
            [
                '序号',
                '区域',
                '对象名称',
                '对象类型',
                '查询日期',
                '地址',
                '地址获取方式',
                '地图匹配状态',
                '信息来源',
            ],
        )
        self.assertEqual(
            abnormal_headers(spec),
            [
                '序号',
                '区域',
                '对象名称',
                '对象类型',
                '异常原因',
                '信息来源',
                '查询日期',
            ],
        )

    def test_column_widths_follow_common_layout(self):
        spec = build_spec()
        self.assertEqual(
            main_column_widths(spec),
            [8, 14, 36, 20, 14, 42, 14, 16, 50],
        )
        self.assertEqual(
            abnormal_column_widths(spec),
            [8, 14, 36, 20, 50, 50, 14],
        )

    def test_rejects_mismatched_domain_headers_and_widths(self):
        with self.assertRaisesRegex(ValueError, '数量必须一致'):
            ResultWorkbookSpec(('区域',), (14, 20))


class MainWorksheetTests(unittest.TestCase):
    def test_main_rows_contain_sequence_date_address_columns_and_hyperlinks(self):
        spec = build_spec()
        official_record = build_address_record(source='official')
        map_record = build_address_record(
            final_address='地图大道88号',
            source='map',
            map_status='poi_match',
        )
        workbook = create_result_workbook(
            spec,
            '结果表',
            '异常表',
            [
                (
                    '结果表',
                    [
                        (['示例区', '示例甲', '甲类'], official_record),
                        (['示例区', '示例乙', '乙类'], map_record),
                    ],
                )
            ],
            [],
        )
        try:
            sheet = workbook['结果表']
            self.assertEqual(
                [sheet.cell(2, column).value for column in range(1, 10)],
                [
                    1,
                    '示例区',
                    '示例甲',
                    '甲类',
                    QUERY_DATE,
                    '示例大道1号',
                    '政府资料',
                    '一致',
                    'https://www.example.gov.cn/list.html | 名录.xlsx | row 3',
                ],
            )
            self.assertEqual(
                [sheet.cell(3, column).value for column in range(1, 10)],
                [
                    2,
                    '示例区',
                    '示例乙',
                    '乙类',
                    QUERY_DATE,
                    '地图大道88号',
                    '地图信息',
                    'POI匹配',
                    'https://www.example.gov.cn/list.html | 名录.xlsx | row 3',
                ],
            )
            source_cell = sheet.cell(2, 9)
            self.assertTrue(source_cell.hyperlink)
            self.assertEqual(
                source_cell.hyperlink.target,
                'https://www.example.gov.cn/list.html',
            )
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet.freeze_panes, 'A2')
        finally:
            workbook.close()

    def test_defaults_query_date_to_today(self):
        spec = build_spec(query_date='')
        workbook = create_result_workbook(
            spec,
            '结果表',
            '异常表',
            [('结果表', [(['示例区', '示例甲', '甲类'], build_address_record())])],
            [],
        )
        try:
            self.assertEqual(
                workbook['结果表'].cell(2, 5).value,
                date.today().isoformat(),
            )
        finally:
            workbook.close()


class AbnormalWorksheetTests(unittest.TestCase):
    def test_abnormal_rows_contain_reason_source_and_date(self):
        spec = build_spec()
        record = build_address_record(
            source_reference='https://www.example.gov.cn/source.html | 名录.xlsx | row 5'
        )
        workbook = create_result_workbook(
            spec,
            '结果表',
            '异常表',
            [('结果表', [])],
            [(['示例区', '示例丁', '丁类'], record, '地址无法定位')],
        )
        try:
            sheet = workbook['异常表']
            self.assertEqual(
                [sheet.cell(2, column).value for column in range(1, 8)],
                [
                    1,
                    '示例区',
                    '示例丁',
                    '丁类',
                    '地址无法定位',
                    'https://www.example.gov.cn/source.html | 名录.xlsx | row 5',
                    QUERY_DATE,
                ],
            )
            source_cell = sheet.cell(2, 6)
            self.assertTrue(source_cell.hyperlink)
        finally:
            workbook.close()


class WorkbookRoundTripTests(unittest.TestCase):
    def test_atomic_write_verifies_multiple_main_sheets_and_abnormal_sheet(self):
        spec = build_spec()
        first_rows = [
            (['一区', '甲对象', '甲类'], build_address_record()),
        ]
        second_rows = [
            (['二区', '乙对象', '乙类'], build_address_record(source='map')),
        ]
        abnormal_rows = [
            (['二区', '丙对象', '丙类'], build_address_record(), '缺少门牌号'),
        ]
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / '结果.xlsx'
            written_path = write_result_workbook_atomically(
                spec,
                '一区',
                '异常表',
                [('一区', first_rows), ('二区', second_rows)],
                abnormal_rows,
                output_path,
            )
            self.assertEqual(written_path, output_path.resolve())
            verify_result_workbook(
                written_path,
                spec,
                [('一区', len(first_rows)), ('二区', len(second_rows))],
                '异常表',
                len(abnormal_rows),
            )
            loaded = load_workbook(written_path, data_only=False)
            try:
                self.assertEqual(loaded.sheetnames, ['一区', '二区', '异常表'])
            finally:
                loaded.close()

    def test_verify_rejects_wrong_headers_and_counts(self):
        spec = build_spec()
        rows = [(['一区', '甲对象', '甲类'], build_address_record())]
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / '结果.xlsx'
            write_result_workbook_atomically(
                spec,
                '一区',
                '异常表',
                [('一区', rows)],
                [],
                output_path,
            )
            with self.assertRaisesRegex(ValueError, '记录数不正确'):
                verify_result_workbook(
                    output_path,
                    spec,
                    [('一区', len(rows) + 1)],
                    '异常表',
                    0,
                )
            wrong_header_path = Path(temporary_directory) / '错误表头.xlsx'
            loaded = load_workbook(output_path, data_only=False)
            loaded['一区']['A1'] = '错误字段'
            loaded.save(wrong_header_path)
            loaded.close()
            with self.assertRaisesRegex(ValueError, '字段不正确'):
                verify_result_workbook(
                    wrong_header_path,
                    spec,
                    [('一区', len(rows))],
                    '异常表',
                    0,
                )


class SourceVocabularyTests(unittest.TestCase):
    def test_core_module_has_no_scenario_vocabulary(self):
        source_path = (
            Path(__file__).resolve().parents[1]
            / 'src'
            / 'query_city_core'
            / 'excel_output.py'
        )
        source_text = source_path.read_text(encoding='utf-8').lower()
        forbidden_terms = (
            'school',
            'university',
            'medical',
            'hospital',
            'institution',
        )
        self.assertFalse(
            [term for term in forbidden_terms if term in source_text]
        )


if __name__ == '__main__':
    unittest.main()
