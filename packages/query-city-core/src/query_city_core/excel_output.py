"""统一地址结果工作簿生成器。

核心只负责通用表格布局、固定输出列、原子保存与自动复核；
领域列名称与含义由调用方通过 ``ResultWorkbookSpec`` 提供。
"""

from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .address.common import (
    MAP_MATCH_STATUS_VALUES,
    format_map_match_status,
)
from .excel_style import (
    ADDRESS_OUTPUT_WIDTHS,
    DATE_CELL_PATTERN,
    add_source_hyperlinks,
    build_address_output_values,
    format_source_reference,
    populate_table_worksheet,
    write_workbook_atomically,
)


__all__ = (
    'ResultWorkbookSpec',
    'abnormal_column_widths',
    'abnormal_headers',
    'create_result_workbook',
    'main_column_widths',
    'main_headers',
    'verify_result_workbook',
    'write_result_workbook_atomically',
)


@dataclass(frozen=True)
class ResultWorkbookSpec:
    """结果工作簿的领域列与固定输出参数。

    domain_headers/domain_widths 是一组配套的领域列；它们会出现在
    序号之后、公共结果列之前。查询日期为空时使用运行当天日期。
    """

    domain_headers: tuple[str, ...]
    domain_widths: tuple[int, ...]
    official_source_label: str = '政府资料'
    map_source_label: str = '地图信息'
    query_date: str = ''

    def __post_init__(self):
        object.__setattr__(self, 'domain_headers', tuple(self.domain_headers))
        object.__setattr__(self, 'domain_widths', tuple(self.domain_widths))
        if len(self.domain_headers) != len(self.domain_widths):
            raise ValueError('领域列数量与领域列宽数量必须一致')


def _resolved_query_date(spec):
    return str(spec.query_date or '').strip() or date.today().isoformat()


def main_headers(spec):
    """返回主表的固定表头顺序。"""
    return [
        '序号',
        *spec.domain_headers,
        '查询日期',
        '地址',
        '地址获取方式',
        '地图匹配状态',
        '信息来源',
    ]


def abnormal_headers(spec):
    """返回异常表的固定表头顺序。"""
    return [
        '序号',
        *spec.domain_headers,
        '异常原因',
        '信息来源',
        '查询日期',
    ]


def main_column_widths(spec):
    """返回主表的列宽，公共地址列沿用统一格式宽度。"""
    return [8, *spec.domain_widths, 14, *ADDRESS_OUTPUT_WIDTHS]


def abnormal_column_widths(spec):
    """返回异常表的列宽。"""
    return [8, *spec.domain_widths, 50, 50, 14]


def _validate_domain_row(domain_values, expected_count, sheet_name):
    if not isinstance(domain_values, (list, tuple)):
        raise ValueError(f'{sheet_name}的每一行领域值必须是列表或元组')
    if len(domain_values) != expected_count:
        raise ValueError(
            f'{sheet_name}的领域列数量必须与表头领域列数量一致'
        )


def _populate_main_sheet(
    workbook,
    sheet_name,
    spec,
    rows,
    *,
    use_active=False,
):
    """向一张主表写入查询日期、地址输出列与样式。"""
    headers = main_headers(spec)
    query_date = _resolved_query_date(spec)
    output_rows = []
    for index, (domain_values, address_record) in enumerate(rows, start=1):
        _validate_domain_row(domain_values, len(spec.domain_headers), sheet_name)
        if not isinstance(address_record, dict):
            raise ValueError(f'{sheet_name}的地址记录必须是对象')
        output_rows.append([
            index,
            *domain_values,
            query_date,
            *build_address_output_values(
                address_record,
                spec.official_source_label,
                spec.map_source_label,
            ),
        ])
    worksheet = workbook.active if use_active else workbook.create_sheet()
    populate_table_worksheet(
        worksheet,
        sheet_name,
        headers,
        output_rows,
        main_column_widths(spec),
    )
    add_source_hyperlinks(worksheet, len(headers))
    return worksheet


def _populate_abnormal_sheet(workbook, sheet_name, spec, rows):
    """写入异常表并应用统一样式。"""
    headers = abnormal_headers(spec)
    query_date = _resolved_query_date(spec)
    output_rows = []
    for index, (domain_values, address_record, reason) in enumerate(
        rows, start=1
    ):
        _validate_domain_row(domain_values, len(spec.domain_headers), sheet_name)
        if not isinstance(address_record, dict):
            raise ValueError(f'{sheet_name}的地址记录必须是对象')
        output_rows.append([
            index,
            *domain_values,
            str(reason or '').strip(),
            format_source_reference(address_record.get('source_reference')),
            query_date,
        ])
    worksheet = workbook.create_sheet()
    populate_table_worksheet(
        worksheet,
        sheet_name,
        headers,
        output_rows,
        abnormal_column_widths(spec),
    )
    add_source_hyperlinks(worksheet, headers.index('信息来源') + 1)
    return worksheet


def create_result_workbook(
    spec,
    main_sheet_name,
    abnormal_sheet_name,
    main_sheets,
    abnormal_rows,
):
    """按固定列顺序创建包含多张主表与一张异常表的工作簿。"""
    if not main_sheets:
        raise ValueError('工作簿至少需要一张主表')
    if not str(main_sheet_name or '').strip():
        raise ValueError('主表名称不能为空')
    if not str(abnormal_sheet_name or '').strip():
        raise ValueError('异常表名称不能为空')
    if main_sheet_name == abnormal_sheet_name:
        raise ValueError('主表名称与异常表名称不能相同')
    seen_main_names = set()
    for sheet_name, rows in main_sheets:
        if not str(sheet_name or '').strip():
            raise ValueError('主表名称不能为空')
        if sheet_name in seen_main_names:
            raise ValueError(f'主表名称重复：{sheet_name}')
        seen_main_names.add(sheet_name)

    workbook = Workbook()
    for index, (sheet_name, rows) in enumerate(main_sheets):
        _populate_main_sheet(
            workbook,
            sheet_name,
            spec,
            rows,
            use_active=index == 0,
        )
    _populate_abnormal_sheet(
        workbook, abnormal_sheet_name, spec, abnormal_rows
    )
    return workbook


def _expected_headers_and_counts(worksheet, expected_headers, expected_count):
    headers = list(expected_headers)
    if [cell.value for cell in worksheet[1]] != headers:
        raise ValueError(f'{worksheet.title}工作表字段不正确')
    if worksheet.max_row - 1 != expected_count:
        raise ValueError(f'{worksheet.title}工作表记录数不正确')
    return headers


def _validate_date_value(worksheet, value, label):
    if not DATE_CELL_PATTERN.fullmatch(str(value or '').strip()):
        raise ValueError(f'{worksheet.title}的{label}必须为 YYYY-MM-DD')


def _validate_source_cell(worksheet, cell):
    value = str(cell.value or '').strip()
    if not value:
        raise ValueError(f'{worksheet.title}存在空信息来源')
    if value.startswith(('http://', 'https://')) and not cell.hyperlink:
        raise ValueError(f'{worksheet.title}的信息来源没有可点击链接')


def _validate_main_worksheet(worksheet, spec, expected_count):
    headers = _expected_headers_and_counts(
        worksheet, main_headers(spec), expected_count
    )
    column_indexes = {header: index for index, header in enumerate(headers, 1)}
    date_column = column_indexes['查询日期']
    address_column = column_indexes['地址']
    acquisition_column = column_indexes['地址获取方式']
    match_column = column_indexes['地图匹配状态']
    source_column = column_indexes['信息来源']
    valid_acquisitions = {spec.official_source_label, spec.map_source_label}
    valid_matches = {
        format_map_match_status(value)
        for value in MAP_MATCH_STATUS_VALUES
    }
    for sequence, row_index in enumerate(
        range(2, worksheet.max_row + 1), start=1
    ):
        if worksheet.cell(row_index, 1).value != sequence:
            raise ValueError(f'{worksheet.title}工作表序号不连续')
        _validate_date_value(
            worksheet,
            worksheet.cell(row_index, date_column).value,
            '查询日期',
        )
        address = str(
            worksheet.cell(row_index, address_column).value or ''
        ).strip()
        if not address:
            raise ValueError(f'{worksheet.title}存在空地址')
        acquisition = str(
            worksheet.cell(row_index, acquisition_column).value or ''
        ).strip()
        if acquisition not in valid_acquisitions:
            raise ValueError(f'{worksheet.title}存在无效地址获取方式')
        match_value = str(
            worksheet.cell(row_index, match_column).value or ''
        ).strip()
        if match_value not in valid_matches:
            raise ValueError(f'{worksheet.title}存在无效地图匹配状态')
        _validate_source_cell(
            worksheet, worksheet.cell(row_index, source_column)
        )


def _validate_abnormal_worksheet(worksheet, spec, expected_count):
    headers = _expected_headers_and_counts(
        worksheet, abnormal_headers(spec), expected_count
    )
    column_indexes = {header: index for index, header in enumerate(headers, 1)}
    reason_column = column_indexes['异常原因']
    source_column = column_indexes['信息来源']
    date_column = column_indexes['查询日期']
    for sequence, row_index in enumerate(
        range(2, worksheet.max_row + 1), start=1
    ):
        if worksheet.cell(row_index, 1).value != sequence:
            raise ValueError(f'{worksheet.title}工作表序号不连续')
        reason = str(
            worksheet.cell(row_index, reason_column).value or ''
        ).strip()
        if not reason:
            raise ValueError(f'{worksheet.title}存在空的异常原因')
        _validate_source_cell(
            worksheet, worksheet.cell(row_index, source_column)
        )
        _validate_date_value(
            worksheet,
            worksheet.cell(row_index, date_column).value,
            '查询日期',
        )


def verify_result_workbook(
    workbook_path,
    spec,
    main_sheets_counts,
    abnormal_sheet_name,
    abnormal_row_count,
):
    """重新打开工作簿并校验表顺序、固定列与内容约束。"""
    loaded = load_workbook(workbook_path, data_only=False)
    try:
        expected_order = [
            sheet_name for sheet_name, _ in main_sheets_counts
        ] + [abnormal_sheet_name]
        if loaded.sheetnames != expected_order:
            raise ValueError('工作簿的工作表名称或顺序不正确')
        for sheet_name, expected_count in main_sheets_counts:
            _validate_main_worksheet(
                loaded[sheet_name], spec, expected_count
            )
        _validate_abnormal_worksheet(
            loaded[abnormal_sheet_name], spec, abnormal_row_count
        )
    finally:
        loaded.close()


def write_result_workbook_atomically(
    spec,
    main_sheet_name,
    abnormal_sheet_name,
    main_sheets,
    abnormal_rows,
    output_path,
):
    """创建、原子保存并自动复核结果工作簿，返回输出路径。"""
    workbook = create_result_workbook(
        spec,
        main_sheet_name,
        abnormal_sheet_name,
        main_sheets,
        abnormal_rows,
    )
    main_sheets_counts = [
        (sheet_name, len(rows)) for sheet_name, rows in main_sheets
    ]

    def validate(temporary_path):
        verify_result_workbook(
            temporary_path,
            spec,
            main_sheets_counts,
            abnormal_sheet_name,
            len(abnormal_rows),
        )

    return write_workbook_atomically(
        workbook,
        Path(output_path),
        validate,
    )
