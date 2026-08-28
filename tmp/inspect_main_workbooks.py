import json
from pathlib import Path

from openpyxl import load_workbook


paths = sorted(
    path
    for path in Path("output").rglob("*.xlsx")
    if path.name == "基础教育学校查询_广州市_2026-08-27.xlsx"
    and ("Deepseek" in str(path) or "Basic_Education" in str(path))
)

result = []
for path in paths:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows, ()))
    values = [list(row) for row in rows]
    result.append(
        {
            "path": str(path),
            "sheet": sheet.title,
            "row_count": len(values),
            "column_count": len(header),
            "header": header,
            "values": values,
        }
    )

if len(result) != 2:
    raise SystemExit(f"expected 2 workbooks, found {len(result)}")

left, right = result


def build_records(workbook):
    records = {}
    duplicate_keys = []
    for row in workbook["values"]:
        record = dict(zip(workbook["header"], row))
        key = (
            record["行政单位"],
            record["学校名称"],
            record["学校类型"],
        )
        if key in records:
            duplicate_keys.append(key)
        records[key] = record
    return records, duplicate_keys


left_records, left_duplicates = build_records(left)
right_records, right_duplicates = build_records(right)
left_keys = set(left_records)
right_keys = set(right_records)
shared_keys = left_keys & right_keys
compare_columns = left["header"][3:]
changed_records = []
changed_column_counts = {}
for key in sorted(shared_keys):
    changed_columns = {
        column: [left_records[key][column], right_records[key][column]]
        for column in compare_columns
        if left_records[key][column] != right_records[key][column]
    }
    if changed_columns:
        changed_records.append({"key": key, "changed_columns": changed_columns})
        for column in changed_columns:
            changed_column_counts[column] = changed_column_counts.get(column, 0) + 1


def count_by_column(records, column):
    counts = {}
    for record in records.values():
        value = record[column] if record[column] is not None else "(空)"
        counts[value] = counts.get(value, 0) + 1
    return dict(sorted(counts.items()))


def count_by_key_part(keys, index):
    counts = {}
    for key in keys:
        value = key[index]
        counts[value] = counts.get(value, 0) + 1
    return dict(sorted(counts.items()))


summary = {
    "left": {key: value for key, value in left.items() if key != "values"},
    "right": {key: value for key, value in right.items() if key != "values"},
    "shared_record_count": len(shared_keys),
    "left_only_count": len(left_keys - right_keys),
    "right_only_count": len(right_keys - left_keys),
    "changed_shared_record_count": len(changed_records),
    "changed_column_counts": dict(sorted(changed_column_counts.items())),
    "left_only_samples": [left_records[key] for key in sorted(left_keys - right_keys)[:10]],
    "right_only_samples": [right_records[key] for key in sorted(right_keys - left_keys)[:10]],
    "changed_record_samples": changed_records[:10],
    "left_type_counts": count_by_column(left_records, "学校类型"),
    "right_type_counts": count_by_column(right_records, "学校类型"),
    "left_only_by_administrative_unit": count_by_key_part(left_keys - right_keys, 0),
    "right_only_by_administrative_unit": count_by_key_part(right_keys - left_keys, 0),
    "left_duplicate_key_count": len(left_duplicates),
    "right_duplicate_key_count": len(right_duplicates),
}
print(json.dumps(summary, ensure_ascii=False))
